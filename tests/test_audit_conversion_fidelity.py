"""Tests for scripts/audit_conversion_fidelity.py helpers."""
from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

from docx import Document
from lxml import etree

_SPEC = importlib.util.spec_from_file_location(
    "audit_conversion_fidelity",
    Path(__file__).resolve().parents[1] / "scripts" / "audit_conversion_fidelity.py",
)
audit = importlib.util.module_from_spec(_SPEC)
assert _SPEC and _SPEC.loader
_SPEC.loader.exec_module(audit)


def _volume(path: Path, body_line: str) -> None:
    """A minimal Federal Register style volume: Contents block, body, endnotes."""
    doc = Document()
    for line in [
        "Contents",
        "Chapter 1—Preliminary 5",  # TOC entry (trailing page number)
        "1 Short title 5",  # TOC entry
        "The Parliament enacts:",  # first real body paragraph
        body_line,
        "Endnotes",
        "Endnote 1—About the endnotes",
    ]:
        doc.add_paragraph(line)
    doc.save(str(path))


def test_docx_body_paras_keeps_body_from_every_volume(tmp_path):
    # Regression pin: each volume carries its own Contents block and its own
    # endnote apparatus. Slicing the concatenation once finds only volume 1's
    # "Endnotes" heading and discards every later volume. `_docx_body_paras`
    # must slice per volume so a multi-volume Act is compared in full.
    v1 = tmp_path / "act-2001-c1-vol1.docx"
    v2 = tmp_path / "act-2001-c1-vol2.docx"
    _volume(v1, "Body text unique to volume one.")
    _volume(v2, "Body text unique to volume two.")

    body = audit._docx_body_paras([v1, v2])

    assert "Body text unique to volume one." in body
    assert "Body text unique to volume two." in body
    # front/back matter is trimmed on both volumes
    assert not any("Endnote 1" in p for p in body)
    assert "Contents" not in body


def test_docx_paths_prefers_compilation_and_guards_prefix_collision(tmp_path):
    docx_dir = tmp_path / "docx"
    docx_dir.mkdir()
    for name in [
        "fair-work-act-2009-c73-vol1.docx",
        "fair-work-act-2009-c73-vol2.docx",
        "fair-work-act-2009-vol1.docx",  # stale legacy download
        "fair-work-(registered-organisations)-act-2009-vol0.docx",  # different Act
    ]:
        (docx_dir / name).write_bytes(b"")

    paths, mode = audit._docx_paths(tmp_path, "fair-work-act-2009", {"comp_num": "73"})

    assert mode == "comp-vol"
    assert [p.name for p in paths] == [
        "fair-work-act-2009-c73-vol1.docx",
        "fair-work-act-2009-c73-vol2.docx",
    ]


def test_docx_paths_orders_volumes_numerically(tmp_path):
    docx_dir = tmp_path / "docx"
    docx_dir.mkdir()
    for n in (1, 2, 10, 11):
        (docx_dir / f"x-act-1997-c266-vol{n}.docx").write_bytes(b"")

    paths, _ = audit._docx_paths(tmp_path, "x-act-1997", {"comp_num": "266"})

    assert [p.name for p in paths] == [
        "x-act-1997-c266-vol1.docx",
        "x-act-1997-c266-vol2.docx",
        "x-act-1997-c266-vol10.docx",
        "x-act-1997-c266-vol11.docx",
    ]


def test_body_slice_trims_real_apparatus_not_front_matter_endnotes():
    # The compilation front matter lists "Endnotes" among the volumes, followed
    # by "Each volume has its own contents". The real apparatus is a bare
    # "Endnotes" followed by "Endnote 1 ...". Only the latter ends the body.
    paras = [
        "Contents",
        "1 Short title 3",
        "The Parliament enacts:",
        "Volume 1",
        "Schedules",
        "Endnotes",  # front-matter copy: must NOT truncate
        "Each volume has its own contents",
        "1 Short title",
        "This Act may be cited as the Example Act.",
        "Endnotes",  # real apparatus: truncates here
        "Endnote 1—About the endnotes",
        "The endnotes provide information about this compilation.",
    ]

    body = audit._body_slice(paras)

    assert "This Act may be cited as the Example Act." in body
    assert "Each volume has its own contents" in body
    assert not any(p.startswith("Endnote 1") for p in body)
    assert body[-1] == "This Act may be cited as the Example Act."


def test_body_slice_trims_pre_2016_notes_to_the_apparatus():
    paras = [
        "The Parliament enacts:",
        "1 Short title",
        "This Act may be cited as the Example Act 1967.",
        "Notes to the Example Act 1967",
        "Note 1",
        "The Note 1 heading is followed by amendment history.",
    ]

    body = audit._body_slice(paras)

    assert body == [
        "The Parliament enacts:",
        "1 Short title",
        "This Act may be cited as the Example Act 1967.",
    ]


def test_akn_paragraphs_includes_table_cells_and_skips_meta():
    xml = (
        '<akomaNtoso xmlns="http://docs.oasis-open.org/legaldocml/ns/akn/3.0">'
        "<act>"
        "<meta><identification><FRBRWork>"
        "<p>meta paragraph excluded</p>"
        "</FRBRWork></identification></meta>"
        "<body>"
        "<heading>Schedule 1  Rates</heading>"
        "<p>The rate for each item is set out in the table.</p>"
        "<table><tr>"
        "<td><p>Item 1</p></td><td><p>5%</p></td>"
        "</tr></table>"
        "</body>"
        "</act></akomaNtoso>"
    )
    root = etree.fromstring(xml.encode("utf-8"))

    paras = audit.akn_paragraphs(root)

    assert "meta paragraph excluded" not in paras
    assert "Schedule 1 Rates" in paras
    assert "The rate for each item is set out in the table." in paras
    # cell text is collected once (via the <td>), not duplicated by the nested <p>
    assert paras.count("Item 1") == 1
    assert "5%" in paras


def test_main_skips_unreadable_docx_without_aborting_run(tmp_path, monkeypatch):
    corpus = tmp_path / "corpus"
    (corpus / "xml").mkdir(parents=True)
    (corpus / "docx").mkdir()

    akn = (
        '<akomaNtoso xmlns="http://docs.oasis-open.org/legaldocml/ns/akn/3.0">'
        "<act><body><p>The Parliament enacts as follows.</p></body></act>"
        "</akomaNtoso>"
    )
    for slug in ("good-act-2000", "bad-act-2000"):
        (corpus / "xml" / f"{slug}.xml").write_text(akn, encoding="utf-8")
    doc = Document()
    doc.add_paragraph("The Parliament enacts as follows.")
    doc.save(str(corpus / "docx" / "good-act-2000-vol0.docx"))
    (corpus / "docx" / "bad-act-2000-vol0.docx").write_bytes(b"not a zip file")

    index = {
        "acts": {
            "good-act-2000": {"xml_path": "xml/good-act-2000.xml", "comp_num": None},
            "bad-act-2000": {"xml_path": "xml/bad-act-2000.xml", "comp_num": None},
        }
    }
    (corpus / "index.json").write_text(json.dumps(index), encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["audit", "--corpus-dir", str(corpus)])
    assert audit.main() == 0

    summary = json.loads((corpus / "reports" / "fidelity" / "SUMMARY.json").read_text())
    assert summary["acts_scanned"] == 1
    assert summary["acts_skipped"] == 1
    assert summary["skipped"][0]["slug"] == "bad-act-2000"
    assert summary["skipped"][0]["reason"].startswith("error:")


def test_main_writes_valid_json_with_within_para(tmp_path, monkeypatch):
    # One Act whose DOCX/AKN differ by a single 1:1 operative-word substitution
    # ("must" -> "may"): compare() yields one equal-length replace opcode, so the
    # Divergence carries a populated within_para list. The per-Act JSON write must
    # serialise those WithinParaResult objects rather than raising TypeError.
    corpus = tmp_path / "corpus"
    (corpus / "xml").mkdir(parents=True)
    (corpus / "docx").mkdir()

    slug = "notice-act-2000"
    akn = (
        '<akomaNtoso xmlns="http://docs.oasis-open.org/legaldocml/ns/akn/3.0">'
        "<act><body><p>The applicant may give the Secretary written notice "
        "within fourteen days after the original notice.</p></body></act>"
        "</akomaNtoso>"
    )
    (corpus / "xml" / f"{slug}.xml").write_text(akn, encoding="utf-8")
    doc = Document()
    doc.add_paragraph(
        "The applicant must give the Secretary written notice "
        "within fourteen days after the original notice."
    )
    doc.save(str(corpus / "docx" / f"{slug}-vol0.docx"))

    index = {"acts": {slug: {"xml_path": f"xml/{slug}.xml", "comp_num": None}}}
    (corpus / "index.json").write_text(json.dumps(index), encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["audit", "--corpus-dir", str(corpus)])
    assert audit.main() == 0

    report = json.loads(
        (corpus / "reports" / "fidelity" / f"{slug}.json").read_text()
    )
    divs = report["divergences"]
    assert any(d.get("within_para") for d in divs)
    for d in divs:
        for w in d.get("within_para", []):
            assert set(w) >= {
                "kind",
                "docx_word_tokens",
                "akn_word_tokens",
                "dropped",
                "inserted",
            }

    summary = json.loads((corpus / "reports" / "fidelity" / "SUMMARY.json").read_text())
    assert "within_para" in summary
    assert summary["within_para"]["by_outer_kind"].get("minor", {}).get("wp_garble") == 1
