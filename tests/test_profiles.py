import re
from dataclasses import replace
from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from lexau.builder import AknBuilder
from lexau.docx_reader import iter_paragraphs
from lexau.models import ActMetadata
from lexau.profiles import ParserProfile

AKN = {"a": "http://docs.oasis-open.org/legaldocml/ns/akn/3.0"}

STATE = ParserProfile(
    name="test-state",
    style_map={"State Part": "ActHead 2", "State Section": "ActHead 5", "State Sub": "Body Text"},
    heading_re=re.compile(r'^(Chapter|Part|Division|Subdivision)[\xa0 ]+(\S+?)(?:[\xa0 ]*[—\t][\xa0 ]*(.*))?$'),
    section_re=re.compile(r'^(\w[\w.\-]*?)\.?[ \t]+(.+)$'),
    def_heading_re=re.compile(r'\b(definitions?|terms\s+used)\b', re.IGNORECASE),
    char_def_styles=frozenset({"CharDef"}),
)


def _meta(**kw) -> ActMetadata:
    kw.setdefault("title_id", "t")
    kw.setdefault("name", "Test Act 2020")
    return ActMetadata(comp_id="t", comp_num="1",
                       year=2020, number=5, effective_date=date(2026, 1, 1), **kw)


def _state_doc() -> Document:
    doc = Document()
    for name in ("State Part", "State Section", "State Sub"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    char = doc.styles.add_style("CharDef", WD_STYLE_TYPE.CHARACTER)
    doc.add_paragraph("Part 1 — Preliminary", style="State Part")
    doc.add_paragraph("1.\tShort title", style="State Section")
    doc.add_paragraph("(1)\tThis Act is the Test Act 2020.", style="State Sub")
    doc.add_paragraph("2.\tTerms used", style="State Section")
    p = doc.add_paragraph(style="State Sub")
    p.add_run("widget", style=char)
    p.add_run(" means a thing.")
    return doc


def _build(doc: Document, profile: ParserProfile | None, meta: ActMetadata):
    b = AknBuilder(meta, profile=profile)
    for p in iter_paragraphs(doc, profile):
        b.add(replace(p, volume_index=0))
    return b.build_with_report({}, last_volume_path=None)[0]


def test_default_uri_unchanged():
    assert _meta().frbr_work_uri == "/akn/au/act/2020/5"


def test_locality_uri():
    assert _meta(locality="vic").frbr_work_uri == "/akn/au-vic/act/2020/5"


def test_without_profile_state_styles_yield_no_sections():
    xml = _build(_state_doc(), None, _meta())
    assert len(xml.xpath("//a:section", namespaces=AKN)) == 0


def test_profile_recovers_sections_and_terms():
    xml = _build(_state_doc(), STATE, _meta(locality="nsw"))
    assert len(xml.xpath("//a:section", namespaces=AKN)) == 2
    assert [t.text for t in xml.xpath("//a:term", namespaces=AKN)] == ["widget"]


def test_empty_profile_equals_none():
    doc = _state_doc()
    a = _build(doc, None, _meta())
    b = _build(doc, ParserProfile(), _meta())
    from lxml import etree
    assert etree.tostring(a) == etree.tostring(b)


def test_safe_name_default_unchanged_and_locality_namespaced():
    assert _meta().safe_name == "test-act-2020"
    assert _meta(locality="vic").safe_name == "vic-test-act-2020"


def test_corpus_roundtrips_locality(tmp_path):
    from lexau.corpus import Corpus
    from lxml import etree

    c = Corpus(tmp_path)
    c.save(_meta(locality="wa", title_id="wa-1"), etree.Element("x"))
    c.save(_meta(name="Other Act 2020"), etree.Element("y"))
    by_name = {m.safe_name: m for m in c.all_metadata()}
    assert by_name["wa-test-act-2020"].locality == "wa"
    assert by_name["other-act-2020"].locality is None
    import json
    idx = json.loads((tmp_path / "index.json").read_text())
    assert "locality" not in idx["acts"]["other-act-2020"]


def test_same_name_different_locality_is_refused(tmp_path):
    import pytest
    from lexau.corpus import Corpus, NameCollisionError
    from lxml import etree

    c = Corpus(tmp_path)
    c.save(_meta(locality="vic", title_id="vic-1"), etree.Element("x"))
    with pytest.raises(NameCollisionError):
        c.save(_meta(title_id="cth-1"), etree.Element("y"))
    with pytest.raises(NameCollisionError):
        c.save(_meta(locality="nsw", title_id="nsw-1"), etree.Element("z"))
    c.save(_meta(locality="nsw", title_id="nsw-1"), etree.Element("z"), allow_name_collision=True)
    assert len(c.all_metadata()) == 2


def test_resave_same_instrument_is_not_a_collision(tmp_path):
    from lexau.corpus import Corpus
    from lxml import etree

    c = Corpus(tmp_path)
    c.save(_meta(locality="vic", title_id="vic-1"), etree.Element("x"))
    c.save(_meta(locality="vic", title_id="vic-1"), etree.Element("x"))
    assert len(c.all_metadata()) == 1
