"""ParserProfile.schedule_clause_res: principle clauses in 'Normal'-style schedule text."""
import re
from dataclasses import replace
from datetime import date

from docx import Document

from lexau.builder import AknBuilder
from lexau.docx_reader import iter_paragraphs
from lexau.models import ActMetadata
from lexau.profiles import ParserProfile

AKN = {"a": "http://docs.oasis-open.org/legaldocml/ns/akn/3.0"}

PROFILE = ParserProfile(
    name="t",
    schedule_heading_re=re.compile(r"^Schedule[\xa0 ]+\d+[\xa0 ]*[\u2014\u2013\t]"),
    schedule_clause_res=(
        re.compile(r"^(?=IPP[\xa0 ]*(?P<number>\d+)[\s:\u2014-])(?P<heading>.+)$", re.DOTALL),
        re.compile(r"^(?P<number>\d+\.\d+)\s+(?P<heading>.+)$", re.DOTALL),
    ),
)


def _build(doc: Document, profile: ParserProfile | None):
    meta = ActMetadata(name="T Act 2000", title_id="t", comp_id="t", comp_num="1", year=2000, number=1,
                       effective_date=date(2026, 1, 1))
    b = AknBuilder(meta, profile=profile)
    for p in iter_paragraphs(doc, profile):
        b.add(replace(p, volume_index=0))
    return b.build_with_report({}, last_volume_path=None)[0]


def _doc() -> Document:
    d = Document()
    d.add_paragraph("Schedule 1\tInformation Privacy Principles")
    d.add_paragraph("IPP 1\tCollection")
    d.add_paragraph("1.1\tAn organisation must not collect personal information unless necessary.")
    d.add_paragraph("1.2\tIt must collect lawfully.")
    d.add_paragraph("IPP 2\tUse and disclosure")
    d.add_paragraph("2.1\tUse only for the primary purpose.")
    return d


def test_schedule_clauses_recovered_with_profile() -> None:
    xml = _build(_doc(), PROFILE)
    sched = xml.xpath("//a:hcontainer[@name='schedule']", namespaces=AKN)
    assert len(sched) == 1
    text = " ".join("".join(sched[0].itertext()).split())
    assert "IPP 1 Collection" in text and "IPP 2 Use and disclosure" in text
    assert "An organisation must not collect personal information unless necessary." in text


def test_no_profile_leaves_schedule_text_as_body() -> None:
    plain = ParserProfile(schedule_heading_re=PROFILE.schedule_heading_re)
    with_clauses = _build(_doc(), PROFILE)
    without = _build(_doc(), plain)
    n = lambda x: len(x.xpath("//a:hcontainer[@name='schedule']//a:hcontainer", namespaces=AKN))
    assert n(with_clauses) > n(without)
