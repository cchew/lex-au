"""ParserProfile.unnumbered_schedule_re and skip_styles: the ACT 'Dictionary' pattern."""
import re
from dataclasses import replace
from datetime import date

from docx import Document

from lexau.builder import AknBuilder
from lexau.docx_reader import iter_paragraphs
from lexau.models import ActMetadata
from lexau.profiles import ParserProfile

AKN = {"a": "http://docs.oasis-open.org/legaldocml/ns/akn/3.0"}

PROFILE = ParserProfile(
    name="t",
    style_map={"Sec": "ActHead 5", "Body": "Body Text"},
    unnumbered_schedule_re=re.compile(r"^Dictionary$"),
    skip_styles=frozenset({"Endnote"}),
    def_heading_re=re.compile(r"\bdictionary\b", re.I),
    definitions_in_schedules=True,
    extended_definitions=True,
)


def _doc() -> Document:
    d = Document()
    for n in ("Sec", "Body", "Endnote"):
        d.styles.add_style(n, 1)
    d.add_paragraph("1\tName of Act", style="Sec")
    d.add_paragraph("This Act is the Test Act 2000.", style="Body")
    d.add_paragraph("Dictionary", style="Body")
    d.add_paragraph("widget—see section 1.", style="Body")
    d.add_paragraph("Endnotes", style="Endnote")
    d.add_paragraph("Amendment history that is not part of the Act.", style="Endnote")
    return d


def _build(profile):
    meta = ActMetadata(name="T Act 2000", title_id="t", comp_id="t", comp_num="1", year=2000, number=1,
                       effective_date=date(2026, 1, 1))
    b = AknBuilder(meta, profile=profile)
    for p in iter_paragraphs(_doc(), profile):
        b.add(replace(p, volume_index=0))
    return b.build_with_report({}, last_volume_path=None)[0]


def test_dictionary_becomes_a_schedule_container_and_endnotes_are_dropped() -> None:
    xml = _build(PROFILE)
    sched = xml.xpath("//a:hcontainer[@name='schedule']", namespaces=AKN)
    assert len(sched) == 1
    assert "".join(sched[0].xpath("a:heading//text()", namespaces=AKN)) == "Dictionary"
    assert "widget—see section 1." in " ".join("".join(sched[0].itertext()).split())
    assert "Amendment history" not in "".join(xml.itertext())


def test_without_the_options_dictionary_stays_body_text() -> None:
    plain = ParserProfile(style_map=PROFILE.style_map)
    xml = _build(plain)
    assert xml.xpath("//a:hcontainer[@name='schedule']", namespaces=AKN) == []
    assert "Amendment history" in "".join(xml.itertext())
