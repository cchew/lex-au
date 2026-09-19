import pytest
from lxml import etree
from datetime import date
from datetime import date as _date
from pathlib import Path
from unittest.mock import patch, MagicMock
from lexau.models import ActMetadata
from lexau.parser import ParsedParagraph, ElementType
from lexau.builder import AknBuilder, inject_lifecycle
from lexau.endnote_parser import AmendmentEvent, EndnoteResult

AKN_NS = "http://docs.oasis-open.org/legaldocml/ns/akn/3.0"

@pytest.fixture
def meta(privacy_meta):
    return privacy_meta

def build_xml(meta, paragraphs):
    b = AknBuilder(meta)
    for p in paragraphs:
        b.add(p)
    return b.build()

def test_root_element_is_akoma_ntoso(meta):
    xml, _ = build_xml(meta, [])
    assert xml.tag == f"{{{AKN_NS}}}akomaNtoso"

def test_frbr_work_uri_in_meta(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    this_elem = xml.find(".//akn:FRBRWork/akn:FRBRthis", ns)
    assert this_elem is not None
    assert "/akn/au/act/1988/119" in this_elem.get("value", "")

def test_single_section_produces_section_element(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    section = xml.find(".//akn:section", ns)
    assert section is not None
    assert section.get("eId") == "sec-1"
    num = section.find("akn:num", ns)
    assert num is not None and num.text == "1"
    heading = section.find("akn:heading", ns)
    assert heading is not None and heading.text == "Short title"

def test_body_text_in_content_p(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    p = xml.find(".//akn:section/akn:content/akn:p", ns)
    assert p is not None
    assert p.text == "This Act is the Privacy Act 1988."

def test_part_contains_sections(meta):
    paragraphs = [
        ParsedParagraph(ElementType.PART, number="I", heading="Preliminary"),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.SECTION, number="2", heading="Commencement"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    part = xml.find(".//akn:part", ns)
    assert part is not None
    assert part.get("eId") == "part-I"
    sections = part.findall("akn:section", ns)
    assert len(sections) == 2
    assert sections[0].get("eId") == "part-I__sec-1"

def test_valid_xml_serialises(meta):
    paragraphs = [
        ParsedParagraph(ElementType.PART, number="I", heading="Preliminary"),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    raw = etree.tostring(xml, encoding="unicode", xml_declaration=False)
    assert "akomaNtoso" in raw
    assert 'eId="part-I"' in raw
    assert 'eId="part-I__sec-1"' in raw


def test_subsection_eid(meta):
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="16", heading="Notification"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="The entity must notify..."),
    ])
    ns = {"akn": AKN_NS}
    subsec = xml.find(".//akn:subsection", ns)
    assert subsec is not None
    assert subsec.get("eId") == "sec-16__subsec-1"


def test_full_hierarchy_eid(meta):
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text=""),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="means..."),
        ParsedParagraph(ElementType.SUBPARAGRAPH, number="i", text="first thing"),
    ])
    ns = {"akn": AKN_NS}
    subpara = xml.find(".//akn:subparagraph", ns)
    assert subpara is not None
    assert subpara.get("eId") == "sec-6__subsec-1__para-a__subpara-i"


def test_paragraph_l_as_subparagraph_inside_paragraph(meta):
    # (l) when the stack has an open PARAGRAPH -> reclassified to SUBPARAGRAPH
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="5", heading="Obligations"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text=""),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text=""),
        ParsedParagraph(ElementType.PARAGRAPH, number="l", text="the ambiguous item"),
    ])
    ns = {"akn": AKN_NS}
    subpara = xml.find(".//akn:subparagraph", ns)
    assert subpara is not None
    assert subpara.get("eId") == "sec-5__subsec-1__para-a__subpara-l"


def test_paragraph_l_as_sibling_paragraph(meta):
    # (l) after (a) at the same depth: (a) is closed off stack before (l) arrives
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="5", heading="Obligations"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text=""),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="the first condition:"),
        ParsedParagraph(ElementType.PARAGRAPH, number="l", text="the roman-ambiguous item"),
    ])
    ns = {"akn": AKN_NS}
    paragraphs = xml.findall(".//akn:paragraph", ns)
    assert len(paragraphs) == 2
    assert paragraphs[1].get("eId") == "sec-5__subsec-1__para-l"


def test_frbr_work_date_is_iso(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    work_date = xml.find(".//akn:FRBRWork/akn:FRBRdate", ns)
    assert work_date is not None
    assert work_date.get("date") == "1988-01-01"


def test_frbr_expression_date_is_iso(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    expr_date = xml.find(".//akn:FRBRExpression/akn:FRBRdate", ns)
    assert expr_date is not None
    assert expr_date.get("date") == "2024-01-01"


def test_preface_toc_emitted(meta):
    paragraphs = [
        ParsedParagraph(ElementType.BODY, text="Privacy Act 1988", raw_style="TOC Heading"),
        ParsedParagraph(ElementType.BODY, text="Part I—Preliminary\t1", raw_style="TOC 1"),
        ParsedParagraph(ElementType.BODY, text="1  Short title\t1", raw_style="TOC 2"),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    preface = xml.find(".//akn:preface", ns)
    assert preface is not None
    toc = preface.find("akn:toc", ns)
    assert toc is not None
    toc_items = toc.findall("akn:tocItem", ns)
    assert len(toc_items) == 2  # both TOC 1 and TOC 2 lines become tocItem elements


def test_preface_non_toc_para_emitted_as_p(meta):
    paragraphs = [
        ParsedParagraph(ElementType.BODY, text="About this compilation", raw_style="Normal"),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    preface = xml.find(".//akn:preface", ns)
    assert preface is not None
    p = preface.find("akn:p", ns)
    assert p is not None and p.text == "About this compilation"


def test_schedule_emitted_as_attachment(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Australian Privacy Principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="APP 1  Open and transparent management"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    attachments = xml.find(".//akn:attachments", ns)
    assert attachments is not None
    hcontainer = attachments.find(".//akn:hcontainer", ns)
    assert hcontainer is not None
    assert hcontainer.get("name") == "schedule"
    assert hcontainer.get("eId") == "schedule-1"


def test_multiple_schedules(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—First Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content of first schedule."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Second Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content of second schedule."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    hcontainers = xml.findall(".//akn:hcontainer[@name='schedule']", ns)
    assert len(hcontainers) == 2
    assert hcontainers[0].get("eId") == "schedule-1"
    assert hcontainers[1].get("eId") == "schedule-2"


def test_body_outside_schedule_not_in_attachments(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Section body text."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    assert xml.find(".//akn:attachments", ns) is None


def test_multi_volume_schedule_before_body_splits_correctly(meta):
    # Mirrors Superannuation Industry (Supervision) Regulations 1994's real
    # shape: Volume 0 is entirely Schedule content, Volume 1 is the real
    # body. Without volume-scoping, the Volume-0 schedule heading would
    # permanently flip _split_stream into schedule mode, silently
    # swallowing Volume 1's section into schedule_paras instead of body.
    paragraphs = [
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—First Schedule", raw_style="ActHead 1", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Content of first schedule.", volume_index=0),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="This Regulation is the Test Regulations 1994.", volume_index=1),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    section = xml.find(".//akn:body/akn:section", ns)
    assert section is not None
    assert section.get("eId") == "sec-1"
    hcontainer = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert hcontainer is not None
    assert hcontainer.get("eId") == "schedule-1"


def test_single_volume_schedule_after_body_unaffected_by_volume_scoping(meta):
    # Regression guard: every paragraph defaults to volume_index=0, so a
    # single-volume Act's existing schedule-after-body shape must produce
    # byte-identical classification to the pre-volume-scoping behaviour.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Australian Privacy Principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="APP 1  Open and transparent management"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    section = xml.find(".//akn:body/akn:section", ns)
    assert section is not None
    hcontainer = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert hcontainer is not None
    assert hcontainer.get("eId") == "schedule-1"


def test_multi_volume_schedule_continues_across_boundary_without_repeated_heading(meta):
    # Mirrors Customs Tariff Act 1995's real shape: volume 0 is body then a
    # Schedule heading near the end, volume 1 is pure continuation of that
    # same schedule with NO repeated "Schedule N" heading at its start.
    # Searching each volume independently for a schedule heading finds none
    # in volume 1 and silently promotes the whole continuation into <body>.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="This Act is the Test Tariff Act 1995.", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa03—Classification of goods", raw_style="ActHead 1", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Chapter 1—Live animals", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="0101.21.00 Pure-bred breeding horses", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="0102.29.00 Other live bovine animals", volume_index=1),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    body = xml.find(".//akn:body", ns)
    assert body is not None
    body_text = "".join(body.itertext())
    assert "Test Tariff Act 1995" in body_text
    assert "Pure-bred breeding horses" not in body_text
    assert "Other live bovine animals" not in body_text

    hcontainers = xml.findall(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert len(hcontainers) == 1
    # eIds are sequential over schedule groups, not derived from the schedule number
    assert hcontainers[0].get("eId") == "schedule-1"
    schedule_text = "".join(hcontainers[0].itertext())
    assert "Pure-bred breeding horses" in schedule_text
    assert "Other live bovine animals" in schedule_text


def test_multi_volume_new_schedule_heading_in_later_volume_starts_new_group(meta):
    # The continuation carry-over must not swallow a genuine new schedule
    # heading that does appear at/after a later volume's start.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="This Act is the Test Act 1995.", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—First Schedule", raw_style="ActHead 1", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Content of first schedule.", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Second Schedule", raw_style="ActHead 1", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="Content of second schedule.", volume_index=1),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    hcontainers = xml.findall(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert [h.get("eId") for h in hcontainers] == ["schedule-1", "schedule-2"]


def test_multi_volume_body_spanning_volumes_with_no_schedules_stays_body(meta):
    # Guard against over-correcting the continuation carry-over: an Act whose
    # body simply spans several volumes and has no schedules at all (e.g.
    # Income Tax Assessment Act 1997, 12 volumes, zero attachments) must keep
    # every volume in <body>. Carrying schedule mode on `body_seen` alone
    # would swallow every volume after the first into a schedule.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="This Act is the Test Act 1997.", volume_index=0),
        ParsedParagraph(ElementType.SECTION, number="2", heading="Commencement", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="This Act commences on Royal Assent.", volume_index=1),
        ParsedParagraph(ElementType.SECTION, number="3", heading="Objects", volume_index=2),
        ParsedParagraph(ElementType.BODY, text="The objects of this Act are as follows.", volume_index=2),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    sections = xml.findall(".//akn:body/akn:section", ns)
    assert [s.get("eId") for s in sections] == ["sec-1", "sec-2", "sec-3"]
    assert xml.find(".//akn:attachments", ns) is None


def test_multi_volume_body_spanning_volumes_before_schedules_stays_body(meta):
    # Same guard, mixed shape: body spans volumes 0 and 1, schedules only
    # begin in volume 2. Volume 1 must stay body, volume 2 must be schedule.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="This Act is the Test Act 1901.", volume_index=0),
        ParsedParagraph(ElementType.SECTION, number="2", heading="Commencement", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="This Act commences on Royal Assent.", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Prescribed goods", raw_style="ActHead 1", volume_index=2),
        ParsedParagraph(ElementType.BODY, text="Content of the schedule.", volume_index=2),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    sections = xml.findall(".//akn:body/akn:section", ns)
    assert [s.get("eId") for s in sections] == ["sec-1", "sec-2"]
    body_text = "".join(xml.find(".//akn:body", ns).itertext())
    assert "commences on Royal Assent" in body_text
    assert "Content of the schedule" not in body_text
    hcontainers = xml.findall(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert len(hcontainers) == 1
    assert "Content of the schedule" in "".join(hcontainers[0].itertext())


def test_leading_schedule_heading_ends_preface_in_first_volume(meta):
    # Volume 0 opens directly with a Schedule heading (no top-level Part/
    # Division before it) -- mirrors SIS Regs 1994's Schedule 1AAA, whose
    # own internal "Part 1" would otherwise wrongly end the preface first.
    paragraphs = [
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01AAA—First Schedule", raw_style="ActHead 1", volume_index=0),
        ParsedParagraph(ElementType.PART, number="1", heading="Preliminary", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="Content of schedule Part 1.", volume_index=0),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=1),
        ParsedParagraph(ElementType.BODY, text="This Regulation is the Test Regulations 1994.", volume_index=1),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    preface = xml.find(".//akn:preface", ns)
    assert preface is None or not preface.findall(".//akn:p", ns)
    hcontainer = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert hcontainer is not None
    # Deviation from the plan's draft, which asserted eId == "schedule-1AAA":
    # schedule eIds are sequential over schedule groups (_build_attachments),
    # never derived from the schedule number. Assert the heading instead --
    # that is what actually pins "Schedule 1AAA ended the preface".
    assert hcontainer.get("eId") == "schedule-1"
    heading = hcontainer.find("akn:heading", ns)
    assert heading is not None and "First Schedule" in "".join(heading.itertext())


def test_leading_part_before_schedule_heading_still_ends_preface_normally(meta):
    # Belt-and-braces regression guard: if a genuine top-level Part appears
    # in volume 0 BEFORE any schedule heading, the existing (unwidened)
    # behaviour must still win -- this is not a schedule-first Act.
    paragraphs = [
        ParsedParagraph(ElementType.PART, number="I", heading="Preliminary", volume_index=0),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title", volume_index=0),
        ParsedParagraph(ElementType.BODY, text="This Act is the Test Act 1994.", volume_index=0),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    part = xml.find(".//akn:body/akn:part", ns)
    assert part is not None
    assert part.get("eId") == "part-I"


def test_schedule_app_clause_hierarchy(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Australian Privacy Principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="APP 1 — Open and transparent management"),
        ParsedParagraph(ElementType.BODY, text="1.1 The object of this APP is to ensure..."),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="have an up-to-date APP privacy policy"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    assert clause is not None
    assert clause.get("eId") == "schedule-1__clause-1"
    num = clause.find("akn:num", ns)
    assert num is not None and num.text == "1"
    subclause = xml.find(".//akn:hcontainer[@name='subclause']", ns)
    assert subclause is not None
    assert "subclause" in subclause.get("eId", "")
    para = xml.find(".//akn:hcontainer[@name='subclause']//akn:paragraph", ns)
    assert para is not None
    assert "para-a" in para.get("eId", "")


def test_schedule_numeric_clause(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Definitions", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="1  General definitions"),
        ParsedParagraph(ElementType.BODY, text="1.1 In this Schedule..."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    assert clause is not None
    assert clause.get("eId") == "schedule-1__clause-1"


def test_schedule_alphanumeric_clause(meta):
    # TG Regs use clause numbers like "1A", "2B" — _CLAUSE_RE must match \d+[A-Z]?
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Test", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="1A  Alpha clause heading"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    assert clause is not None
    assert clause.get("eId") == "schedule-1__clause-1A"


def test_schedule_section_typed_clause(meta):
    # TG Regs: clause headings parsed as SECTION elements (not BODY text) with num+heading fields
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Essential principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.SECTION, number="7", heading="Chemical properties"),
        ParsedParagraph(ElementType.BODY, text="A device must be designed to minimise risk."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    assert clause is not None
    assert clause.get("eId") == "schedule-1__clause-7"
    heading = clause.find("akn:heading", ns)
    assert heading is not None and heading.text == "Chemical properties"


def test_schedule_section_typed_dotted_subclause(meta):
    # TG Regs: dotted clause numbers (7.1) parsed as SECTION → subclause under current clause
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Essential principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.SECTION, number="7", heading="Chemical properties"),
        ParsedParagraph(ElementType.SECTION, number="7.1", heading="Choice of materials"),
        ParsedParagraph(ElementType.BODY, text="Materials must be biocompatible."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    subclause = xml.find(".//akn:hcontainer[@name='subclause']", ns)
    assert subclause is not None
    assert subclause.get("eId") == "schedule-1__clause-7__subclause-7-1"
    heading = subclause.find("akn:heading", ns)
    assert heading is not None and heading.text == "Choice of materials"


def test_schedule_prose_after_note_stays_under_subclause(meta):
    # C1 regression: a NOTE catch-all element runs first and opens a <content>
    # under the clause; the following BODY prose asked for the subclause. Without
    # parent affinity in _content_for the prose was appended into the clause-level
    # <content> and lost its subclause eId association.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Essential principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.SECTION, number="7", heading="Chemical properties"),
        ParsedParagraph(ElementType.SECTION, number="7.1", heading="Choice of materials"),
        ParsedParagraph(ElementType.NOTE, text="Note: biocompatibility is assessed under clause 8."),
        ParsedParagraph(ElementType.BODY, text="Materials must not degrade in service."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    subclause = xml.find(".//akn:hcontainer[@name='subclause']", ns)
    assert subclause is not None
    assert subclause.get("eId") == "schedule-1__clause-7__subclause-7-1"
    # The trailing prose paragraph must live under the subclause, not the clause.
    sub_ps = subclause.findall("akn:content/akn:p", ns)
    sub_texts = [p.text for p in sub_ps]
    assert "Materials must not degrade in service." in sub_texts
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    clause_ps = clause.findall("akn:content/akn:p", ns)
    assert "Materials must not degrade in service." not in [p.text for p in clause_ps]


def test_schedule_subsection_typed_subclause(meta):
    # TG Regs: SUBSECTION paragraphs with num inside a schedule clause → numbered subclause
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Essential principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.SECTION, number="2", heading="Design principles"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="The manufacturer must adopt safe design solutions."),
        ParsedParagraph(ElementType.SUBSECTION, number="2", text="Without limiting subclause (1), solutions must..."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clause = xml.find(".//akn:hcontainer[@name='clause']", ns)
    assert clause is not None
    assert clause.get("eId") == "schedule-1__clause-2"
    subclauses = clause.findall("akn:hcontainer[@name='subclause']", ns)
    assert len(subclauses) == 2
    assert subclauses[0].get("eId") == "schedule-1__clause-2__subclause-1"
    assert subclauses[1].get("eId") == "schedule-1__clause-2__subclause-2"


def test_build_attachments_returns_tuple(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Test", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="1  First Clause"),
        ParsedParagraph(ElementType.BODY, text="2  Second Clause"),
    ]
    # build() uses _build_attachments internally — just verify it works end-to-end
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    clauses = xml.findall(".//akn:hcontainer[@name='clause']", ns)
    assert len(clauses) == 2


def test_schedule_gazetted_num_emitted_and_ordered(meta):
    # Task 3: emit gazetted schedule `<num>` child in schedule hcontainers,
    # preserving the gazetted number verbatim even when schedule sequence skips.
    # Case 1: numeric skip (1, 2, 5) — third schedule has <num>5</num>, eId=schedule-3
    # Case 2: roman numeral (Schedule IV -> <num>IV</num>)
    # Case 3: malformed number (not matching regex -> no <num>, <heading> unchanged)
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—First Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content 1"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Second Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content 2"),
        # Skip Schedule 3, 4 — jump to 5 (case 1: numeric skip)
        ParsedParagraph(ElementType.BODY, text="Schedule\xa05—Fifth Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content 5"),
        # Roman numeral (case 2)
        ParsedParagraph(ElementType.BODY, text="Schedule\xa0IV—Roman Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content IV"),
        # Malformed heading (case 3: no-match on regex, treated as schedule but no <num>)
        # Note: "Schedule One" (spelled out) doesn't match _SCHEDULE_RE, but if it somehow
        # gets here, the heading should remain unchanged
        ParsedParagraph(ElementType.BODY, text="Schedule One—Spelled Out Schedule", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Spelled out content"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    # Find all schedule hcontainers
    hcontainers = xml.findall(".//akn:hcontainer[@name='schedule']", ns)
    # Note: "Schedule One" won't match _SCHEDULE_RE, so it won't be recognized as a schedule
    # heading. We'll only have 4 schedules: 1, 2, 5, IV
    assert len(hcontainers) == 4, f"Expected 4 schedules, got {len(hcontainers)}"

    # Case 1: First schedule — numeric 1
    assert hcontainers[0].get("eId") == "schedule-1"
    num1 = hcontainers[0].find("akn:num", ns)
    assert num1 is not None
    assert num1.text == "1"
    heading1 = hcontainers[0].find("akn:heading", ns)
    assert heading1 is not None
    # <num> must precede <heading> in XML
    num1_idx = list(hcontainers[0]).index(num1)
    heading1_idx = list(hcontainers[0]).index(heading1)
    assert num1_idx < heading1_idx, "<num> must come before <heading>"

    # Case 1b: Second schedule — numeric 2
    assert hcontainers[1].get("eId") == "schedule-2"
    num2 = hcontainers[1].find("akn:num", ns)
    assert num2 is not None
    assert num2.text == "2"

    # Case 1c: Third schedule — gazetted number is 5, but eId is schedule-3
    assert hcontainers[2].get("eId") == "schedule-3"
    num5 = hcontainers[2].find("akn:num", ns)
    assert num5 is not None
    assert num5.text == "5", "gazetted number should be preserved as-is"
    heading5 = hcontainers[2].find("akn:heading", ns)
    assert heading5 is not None and "Fifth Schedule" in heading5.text
    # Verify element order again
    num5_idx = list(hcontainers[2]).index(num5)
    heading5_idx = list(hcontainers[2]).index(heading5)
    assert num5_idx < heading5_idx, "<num> must come before <heading>"

    # Case 2: Roman numeral — Schedule IV
    assert hcontainers[3].get("eId") == "schedule-4"
    numRom = hcontainers[3].find("akn:num", ns)
    assert numRom is not None
    assert numRom.text == "IV", "roman numeral should be preserved verbatim"
    headingRom = hcontainers[3].find("akn:heading", ns)
    assert headingRom is not None and "Roman Schedule" in headingRom.text


def test_schedule_num_guard_on_non_matching_heading():
    # Task 3, Case 3 (refined): verify the no-match guard handles a schedule heading
    # that doesn't match _SCHEDULE_RE. This tests that if a heading like "Schedule L"
    # (roman numeral outside [IVX] range) somehow makes it to _build_attachments,
    # the code handles it gracefully: no <num> emitted, heading text unchanged.
    from lxml import etree
    from lexau.builder import _build_attachments, AKN_NS

    ns = {"akn": AKN_NS}

    # Manually construct a schedule group with a heading that doesn't match the regex.
    # "Schedule L" (50 in roman numerals) won't match _SCHEDULE_RE because L is not
    # in the [IVX] character class.
    heading_text = "Schedule L—Numbering"
    schedule_group = [
        ParsedParagraph(ElementType.BODY, text=heading_text, raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Content of schedule."),
    ]
    schedule_groups = [schedule_group]

    # Call _build_attachments directly
    attachments_el, clause_count = _build_attachments(schedule_groups)

    assert attachments_el is not None
    hcontainer = attachments_el.find(".//akn:hcontainer[@name='schedule']", ns)
    assert hcontainer is not None, "schedule hcontainer must be created"
    assert hcontainer.get("eId") == "schedule-1"

    # Verify no <num> element (regex didn't match)
    num_el = hcontainer.find("akn:num", ns)
    assert num_el is None, "no <num> should be emitted when heading doesn't match _SCHEDULE_RE"

    # Verify heading is byte-identical to input (no slicing or stripping applied)
    heading_el = hcontainer.find("akn:heading", ns)
    assert heading_el is not None
    assert heading_el.text == heading_text, \
        f"heading should be unchanged when regex doesn't match. Expected {heading_text!r}, got {heading_el.text!r}"


def test_authorial_note_emitted(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="16", heading="Notification"),
        ParsedParagraph(ElementType.NOTE, text="Note: See also section 6."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    note = xml.find(".//akn:authorialNote", ns)
    assert note is not None
    assert note.get("placement") == "end"
    assert note.get("marker") == "1"
    assert note.get("eId") == "note-1"
    p = note.find("akn:content/akn:p", ns)
    assert p is not None
    assert p.text == "Note: See also section 6."


def test_example_emitted(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"),
        ParsedParagraph(ElementType.EXAMPLE, text="Example: A person who transfers data..."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    ex = xml.find(".//akn:hcontainer[@name='example']", ns)
    assert ex is not None
    p = ex.find("akn:content/akn:p", ns)
    assert p is not None
    assert p.text == "Example: A person who transfers data..."


def test_penalty_emitted(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="13G", heading="Offences"),
        ParsedParagraph(ElementType.PENALTY, text="Penalty: 60 penalty units."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    pen = xml.find(".//akn:hcontainer[@name='penalty']", ns)
    assert pen is not None
    p = pen.find("akn:content/akn:p", ns)
    assert p is not None
    assert p.text == "Penalty: 60 penalty units."


def test_table_emitted_as_akn_table(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Fees"),
        ParsedParagraph(
            ElementType.TABLE,
            table_rows=[["Header 1", "Header 2"], ["Data 1", "Data 2"]],
        ),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    table = xml.find(".//akn:table", ns)
    assert table is not None
    rows = table.findall("akn:tr", ns)
    assert len(rows) == 2
    ths = rows[0].findall("akn:th", ns)
    assert len(ths) == 2
    assert ths[0].text == "Header 1"
    tds = rows[1].findall("akn:td", ns)
    assert len(tds) == 2
    assert tds[0].text == "Data 1"


def test_empty_table_rows_skipped(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Fees"),
        ParsedParagraph(ElementType.TABLE, table_rows=[]),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    table = xml.find(".//akn:table", ns)
    assert table is not None
    rows = table.findall("akn:tr", ns)
    assert len(rows) == 0


def test_schedule_table_emitted_as_akn_table(meta):
    # A TABLE ParsedParagraph routed into a schedule group must reach the AKN
    # <schedule> hcontainer as a real <table> with <tr>/<td> and cell text.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa03—Rates of duty", raw_style="ActHead 1"),
        ParsedParagraph(
            ElementType.TABLE,
            table_rows=[
                ["2710.19", "Petroleum oils", "$0.442 per litre"],
                ["2710.20", "Biodiesel", "Free"],
            ],
        ),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    table = xml.find(
        ".//akn:attachments//akn:hcontainer[@name='schedule']//akn:table", ns
    )
    assert table is not None
    rows = table.findall("akn:tr", ns)
    assert len(rows) == 2
    assert [td.text for td in rows[0].findall("akn:td", ns)] == [
        "2710.19",
        "Petroleum oils",
        "$0.442 per litre",
    ]
    assert rows[1].findall("akn:td", ns)[1].text == "Biodiesel"


def test_schedule_headerless_table_renders_all_td(meta):
    # Legislation rate/repeal tables usually have no header row. The schedule
    # branch must emit every row as <td> (no <th>, no <thead>) without crashing.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Repeal of Acts", raw_style="ActHead 1"),
        ParsedParagraph(
            ElementType.TABLE,
            table_rows=[
                ["Customs Act 1901", "No. 6, 1901"],
                ["Excise Act 1901", "No. 9, 1901"],
                ["Sales Tax Act 1930", "No. 25, 1930"],
            ],
        ),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    table = xml.find(
        ".//akn:attachments//akn:hcontainer[@name='schedule']//akn:table", ns
    )
    assert table is not None
    assert table.findall(".//akn:th", ns) == []
    rows = table.findall("akn:tr", ns)
    assert len(rows) == 3
    assert all(len(r.findall("akn:td", ns)) == 2 for r in rows)
    assert rows[0].findall("akn:td", ns)[0].text == "Customs Act 1901"
    assert rows[2].findall("akn:td", ns)[1].text == "No. 25, 1930"


def test_schedule_prose_only_has_no_spurious_table(meta):
    # Regression: a schedule with only prose must be unchanged by the TABLE
    # branch -- no <table> emitted, prose <p> still present.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="This Act is the Privacy Act 1988."),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Australian Privacy Principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="APP 1  Open and transparent management"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    schedule = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert schedule is not None
    assert schedule.find(".//akn:table", ns) is None
    assert b"<table" not in etree.tostring(schedule)
    text = "".join(schedule.itertext())
    assert "Open and transparent management" in text


def test_schedule_prose_after_table_keeps_document_order(meta):
    # prose / TABLE / prose in one schedule context: the trailing prose must
    # open a FRESH <content> that sits AFTER the <table> in document order, not
    # merge into the pre-table <content> (which would reorder it above the table).
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Method", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Use the factor from the table below."),
        ParsedParagraph(ElementType.TABLE, table_rows=[["Age", "Factor"], ["55", "1.20"]]),
        ParsedParagraph(ElementType.BODY, text="Round the result to two decimal places."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    schedule = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert schedule is not None
    kids = [etree.QName(c).localname for c in schedule]
    assert kids == ["num", "heading", "content", "table", "content"]
    contents = schedule.findall("akn:content", ns)
    assert "".join(contents[0].itertext()).strip() == "Use the factor from the table below."
    assert "".join(contents[1].itertext()).strip() == "Round the result to two decimal places."
    table = schedule.find("akn:table", ns)
    after = table.getnext()
    assert after is not None and etree.QName(after).localname == "content"
    assert "Round the result" in "".join(after.itertext())


def test_schedule_empty_table_rows_no_crash(meta):
    # Mirror of the body path's test_empty_table_rows_skipped: an empty
    # table_rows list in a schedule emits a bare <table> and does not crash.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Empty", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.TABLE, table_rows=[]),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    table = xml.find(
        ".//akn:attachments//akn:hcontainer[@name='schedule']//akn:table", ns
    )
    assert table is not None
    assert table.findall("akn:tr", ns) == []


def test_schedule_prose_after_paragraph_keeps_document_order(meta):
    # prose / PARAGRAPH / prose in one schedule context: the trailing prose must
    # open a FRESH <content> that sits AFTER the <paragraph> in document order, not
    # merge into the pre-paragraph <content> (which would reorder it above the paragraph).
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Method", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Use the factor from the table below."),
        ParsedParagraph(ElementType.PARAGRAPH, number="1", text="The paragraph content."),
        ParsedParagraph(ElementType.BODY, text="Round the result to two decimal places."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    schedule = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert schedule is not None
    kids = [etree.QName(c).localname for c in schedule]
    assert kids == ["num", "heading", "content", "paragraph", "content"]
    contents = schedule.findall("akn:content", ns)
    assert len(contents) == 2
    assert "".join(contents[0].itertext()).strip() == "Use the factor from the table below."
    assert "".join(contents[1].itertext()).strip() == "Round the result to two decimal places."
    paragraph = schedule.find("akn:paragraph", ns)
    after = paragraph.getnext()
    assert after is not None and etree.QName(after).localname == "content"
    assert "Round the result" in "".join(after.itertext())


def test_schedule_prose_after_subparagraph_keeps_document_order(meta):
    # prose / SUBPARAGRAPH / prose in one schedule context: the trailing prose must
    # open a FRESH <content> that sits AFTER the <subparagraph> in document order, not
    # merge into the pre-subparagraph <content> (which would reorder it above the subparagraph).
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Method", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Use the factor from the table below."),
        ParsedParagraph(ElementType.SUBPARAGRAPH, number="i", text="The subparagraph content."),
        ParsedParagraph(ElementType.BODY, text="Round the result to two decimal places."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    schedule = xml.find(".//akn:attachments//akn:hcontainer[@name='schedule']", ns)
    assert schedule is not None
    kids = [etree.QName(c).localname for c in schedule]
    assert kids == ["num", "heading", "content", "subparagraph", "content"]
    contents = schedule.findall("akn:content", ns)
    assert len(contents) == 2
    assert "".join(contents[0].itertext()).strip() == "Use the factor from the table below."
    assert "".join(contents[1].itertext()).strip() == "Round the result to two decimal places."
    subpara = schedule.find("akn:subparagraph", ns)
    after = subpara.getnext()
    assert after is not None and etree.QName(after).localname == "content"
    assert "Round the result" in "".join(after.itertext())


def test_level4_emitted_with_lowercase_eid(meta):
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="45", heading="Tests"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text=""),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text=""),
        ParsedParagraph(ElementType.SUBPARAGRAPH, number="i", text=""),
        ParsedParagraph(ElementType.LEVEL4, number="A", text="the entity must comply"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    level4 = xml.find(".//akn:hcontainer[@name='level4']", ns)
    assert level4 is not None
    assert level4.get("eId") == "sec-45__subsec-1__para-a__subpara-i__level4-a"
    num = level4.find("akn:num", ns)
    assert num is not None and num.text == "A"
    p = level4.find("akn:content/akn:p", ns)
    assert p is not None and "comply" in p.text


def test_frbr_work_has_country(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRcountry", ns)
    assert el is not None
    assert el.get("value") == "au"


def test_frbr_work_has_subtype(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRsubtype", ns)
    assert el is not None
    assert el.get("value") == "act"


def test_frbr_work_has_number(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRnumber", ns)
    assert el is not None
    assert el.get("value") == "119"  # privacy_meta.number


def test_frbr_work_has_name(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRname", ns)
    assert el is not None
    assert el.get("value") == "privacy-act-1988"


def test_frbr_work_is_prescriptive(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRprescriptive", ns)
    assert el is not None
    assert el.get("value") == "true"


def test_frbr_work_is_authoritative(meta):
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    el = xml.find(".//akn:FRBRWork/akn:FRBRauthoritative", ns)
    assert el is not None
    assert el.get("value") == "true"


def _meta_with_keywords():
    return ActMetadata(
        name="Privacy Act 1988",
        title_id="C2004A03712",
        comp_id="C2024C00280",
        comp_num="52",
        year=1988,
        number=119,
        effective_date=_date(2024, 1, 1),
        subject_keywords=["Privacy", "Data Protection"],
    )


def test_classification_keywords_in_meta(meta):
    m = _meta_with_keywords()
    b = AknBuilder(m)
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    classification = xml.find(".//akn:meta/akn:classification", ns)
    assert classification is not None
    assert classification.get("source") == "#legislation-gov-au"
    keywords = classification.findall("akn:keyword", ns)
    assert len(keywords) == 2
    show_as_values = {kw.get("showAs") for kw in keywords}
    assert "Privacy" in show_as_values
    assert "Data Protection" in show_as_values
    value_attrs = {kw.get("value") for kw in keywords}
    assert "privacy" in value_attrs
    assert "data-protection" in value_attrs


def test_no_keywords_no_classification(meta):
    # meta fixture has subject_keywords=[] (default)
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    classification = xml.find(".//akn:meta/akn:classification", ns)
    assert classification is None


def test_classification_before_references(meta):
    m = _meta_with_keywords()
    b = AknBuilder(m)
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    meta_el = xml.find(".//akn:meta", ns)
    tags = [c.tag.split("}")[1] for c in meta_el]
    assert tags.index("classification") < tags.index("references"), (
        f"<classification> must precede <references> in <meta>; got order: {tags}"
    )


def _meta_with_long_title():
    m = ActMetadata(
        name="Privacy Act 1988",
        title_id="C2004A03712",
        comp_id="C2024C00280",
        comp_num="52",
        year=1988,
        number=119,
        effective_date=_date(2024, 1, 1),
        long_title="An Act to make provision to protect the privacy of individuals",
    )
    return m


def test_long_title_in_preface(meta):
    m = _meta_with_long_title()
    b = AknBuilder(m)
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    lt = xml.find(".//akn:preface/akn:longTitle/akn:p", ns)
    assert lt is not None
    assert "privacy of individuals" in lt.text


def test_no_long_title_no_long_title_element(meta):
    # meta fixture has long_title="" (default)
    xml, _ = build_xml(meta, [])
    ns = {"akn": AKN_NS}
    lt = xml.find(".//akn:longTitle", ns)
    assert lt is None


def test_enacting_formula_detected(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.BODY, text="The Parliament of Australia enacts:"))
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"))
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    formula = xml.find(".//akn:preface/akn:formula", ns)
    assert formula is not None
    assert formula.get("name") == "enacting"
    assert "enacts" in "".join(formula.itertext())


def test_whereas_recital_detected(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.BODY, text="WHEREAS the Parliament intends to protect privacy:"))
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"))
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    # <preamble> must be a sibling of <preface> under <act>, NOT inside <preface>
    assert xml.find(".//akn:preface/akn:preamble", ns) is None, "<preamble> must NOT be inside <preface>"
    recital = xml.find(".//akn:act/akn:preamble/akn:recitals/akn:recital", ns)
    assert recital is not None


def test_ordinary_preface_para_unaffected(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.BODY, text="This is a normal preface paragraph."))
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"))
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    # Should be a bare <p> not a <formula> or <recital>
    formula = xml.find(".//akn:preface/akn:formula", ns)
    assert formula is None
    p = xml.find(".//akn:preface/akn:p", ns)
    assert p is not None
    assert p.text == "This is a normal preface paragraph."


def test_tlcterm_in_references_after_build_with_report(meta):
    b = AknBuilder(meta)
    # Simulate a Definitions section
    b.add(ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"))
    b.add(ParsedParagraph(ElementType.BODY, text='"personal information" means information about an individual.'))
    b.add(ParsedParagraph(ElementType.SECTION, number="7", heading="Objects"))
    corpus_index = {}
    xml, report = b.build_with_report(corpus_index)
    assert report.terms_found == 1
    ns = {"akn": AKN_NS}
    tlc = xml.find('.//akn:references/akn:TLCTerm[@eId="term-personal-information"]', ns)
    assert tlc is not None
    assert tlc.get("showAs") == "personal information"


def test_refs_injected_inside_def_element(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"))
    b.add(ParsedParagraph(ElementType.BODY,
        text='"notifiable data breach" means a data breach as described in section 26.'))
    corpus_index = {}
    xml, _ = b.build_with_report(corpus_index)
    ns = {"akn": AKN_NS}
    def_el = xml.find(f".//{{{AKN_NS}}}def")
    assert def_el is not None
    # inject_refs should process text within <def> — check for a <ref> inside it
    ref_in_def = def_el.find(f"{{{AKN_NS}}}ref")
    assert ref_in_def is not None
    assert "sec-26" in ref_in_def.get("href", "")


def test_authorial_notes_get_eids(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Objects"))
    b.add(ParsedParagraph(ElementType.NOTE, text="Note: See also section 6."))
    b.add(ParsedParagraph(ElementType.NOTE, text="Note: See also section 7."))
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    notes = xml.findall(".//akn:authorialNote", ns)
    assert len(notes) == 2
    eids = [n.get("eId") for n in notes]
    assert "note-1" in eids
    assert "note-2" in eids


def test_note_ref_injected_for_bracket_marker(meta):
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Objects"))
    b.add(ParsedParagraph(ElementType.BODY, text="This provision applies [note 1]."))
    b.add(ParsedParagraph(ElementType.NOTE, text="Note 1: See section 6."))
    corpus_index = {}
    xml, report = b.build_with_report(corpus_index)
    assert report.note_refs_injected == 1
    ns = {"akn": AKN_NS}
    note_ref = xml.find(".//akn:noteRef[@href='#note-1']", ns)
    assert note_ref is not None
    assert note_ref.get("marker") == "1"


# --- blockList tests ---

def test_blocklist_basic(meta):
    """Two consecutive LIST_ITEM paragraphs at level 0 produce one <blockList> with two <item> children."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="4", heading="Obligations"),
        ParsedParagraph(ElementType.LIST_ITEM, number="0", text="First item"),
        ParsedParagraph(ElementType.LIST_ITEM, number="0", text="Second item"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    blocklists = xml.findall(".//akn:blockList", ns)
    assert len(blocklists) == 1
    items = blocklists[0].findall("akn:item", ns)
    assert len(items) == 2
    assert items[0].get("eId") == "sec-4__list-1__item-1"
    assert items[1].get("eId") == "sec-4__list-1__item-2"


def test_blocklist_level_change(meta):
    """Level 0 item followed by level 1 item produces two separate <blockList> elements."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="5", heading="Requirements"),
        ParsedParagraph(ElementType.LIST_ITEM, number="0", text="Top level item"),
        ParsedParagraph(ElementType.LIST_ITEM, number="1", text="Nested item"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    blocklists = xml.findall(".//akn:blockList", ns)
    assert len(blocklists) == 2
    assert blocklists[0].get("eId") == "sec-5__list-1"
    assert blocklists[1].get("eId") == "sec-5__list-2"


def test_blocklist_flush_on_subsection(meta):
    """LIST_ITEM followed by SUBSECTION: <blockList> is closed, <subsection> is a sibling."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="6", heading="Powers"),
        ParsedParagraph(ElementType.LIST_ITEM, number="0", text="An item in the list"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="A subsection follows"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    section = xml.find(".//akn:section", ns)
    assert section is not None
    # blockList and subsection should both be direct children of the section
    blocklist = section.find("akn:blockList", ns)
    subsection = section.find("akn:subsection", ns)
    assert blocklist is not None
    assert subsection is not None


def test_blocklist_num_extraction(meta):
    """Text '(a) the person is a resident' produces <num>(a)</num><p>the person is a resident</p>."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="7", heading="Conditions"),
        ParsedParagraph(ElementType.LIST_ITEM, number="0", text="(a) the person is a resident"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    item = xml.find(".//akn:blockList/akn:item", ns)
    assert item is not None
    num = item.find("akn:num", ns)
    assert num is not None and num.text == "(a)"
    p = item.find("akn:p", ns)
    assert p is not None and p.text == "the person is a resident"


# --- lifecycle / eventRef tests (Task 5) ---

def _make_events(*args) -> list[AmendmentEvent]:
    """Helper: build AmendmentEvent list from (provision, effect, act_number, act_year) tuples."""
    return [
        AmendmentEvent(provision=prov, effect=eff, act_number=num, act_year=yr)
        for prov, eff, num, yr in args
    ]


def test_lifecycle_emitted(meta):
    """build_with_report with mocked endnote events emits <lifecycle> inside <meta>."""
    events = _make_events(("s 6", "am", 99, 2010))
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, report = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    lifecycle = xml.find(".//akn:meta/akn:lifecycle", ns)
    assert lifecycle is not None
    assert report.amendment_events_parsed == 1


def test_lifecycle_creation_event(meta):
    """<eventRef type='generation' eId='evt-creation'> is always present when lifecycle is emitted."""
    events = _make_events(("s 6", "am", 42, 2005))
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    creation = xml.find(
        ".//akn:lifecycle/akn:eventRef[@eId='evt-creation'][@type='generation']", ns
    )
    assert creation is not None
    assert creation.get("date") == "1988-01-01"  # meta.year = 1988


def test_lifecycle_amendment_events(meta):
    """Two unique amending Acts produce two <eventRef type='amendment'>."""
    events = _make_events(
        ("s 6", "am", 70, 2009),
        ("s 7", "am", 109, 2004),
    )
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    amd_events = xml.findall(
        ".//akn:lifecycle/akn:eventRef[@type='amendment']", ns
    )
    assert len(amd_events) == 2
    sources = {e.get("source") for e in amd_events}
    assert "/akn/au/act/2009/70" in sources
    assert "/akn/au/act/2004/109" in sources


def test_lifecycle_dedup(meta):
    """Same act_number/act_year in multiple rows produces only one <eventRef type='amendment'>."""
    events = _make_events(
        ("s 6", "am", 70, 2009),
        ("s 8", "rep", 70, 2009),  # same Act — should be deduped
        ("s 9", "ad", 70, 2009),   # same Act again
    )
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    amd_events = xml.findall(
        ".//akn:lifecycle/akn:eventRef[@type='amendment']", ns
    )
    assert len(amd_events) == 1
    assert amd_events[0].get("source") == "/akn/au/act/2009/70"


def test_lifecycle_skipped_no_path(meta):
    """last_volume_path=None → no <lifecycle> emitted."""
    b = AknBuilder(meta)
    xml, report = b.build_with_report({}, last_volume_path=None)

    ns = {"akn": AKN_NS}
    lifecycle = xml.find(".//akn:lifecycle", ns)
    assert lifecycle is None
    assert report.amendment_events_parsed == 0


# --- temporalData tests (Task 6) ---

def test_temporal_data_emitted(meta):
    """<temporalData> is present inside <meta> when lifecycle is present."""
    events = _make_events(("s 6", "am", 99, 2010))
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    td = xml.find(".//akn:meta/akn:temporalData", ns)
    assert td is not None


def test_temporal_group_exists(meta):
    """<temporalGroup eId='tg-1'> is a child of <temporalData>."""
    events = _make_events(("s 6", "am", 99, 2010))
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    tg = xml.find(".//akn:temporalData/akn:temporalGroup[@eId='tg-1']", ns)
    assert tg is not None


def test_time_interval_open(meta):
    """<timeInterval start='#evt-creation'> has no end attribute (open-ended)."""
    events = _make_events(("s 6", "am", 99, 2010))
    fake_result = EndnoteResult(amendment_events=events)

    b = AknBuilder(meta)
    with patch("lexau.builder.parse_endnotes", return_value=fake_result), \
         patch("lexau.builder.DocxDocument", MagicMock()):
        xml, _ = b.build_with_report({}, last_volume_path=Path("fake.docx"))

    ns = {"akn": AKN_NS}
    ti = xml.find(".//akn:temporalGroup/akn:timeInterval[@start='#evt-creation']", ns)
    assert ti is not None
    assert ti.get("end") is None


# --- passiveModifications tests (Task 7) ---

def _build_tree_with_section(meta, section_num: str = "6"):
    """Build a minimal AKN tree containing a section with eId sec-{section_num}."""
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number=section_num, heading="Definitions"))
    xml, _ = b.build()
    return xml


def test_passive_mod_emitted(meta):
    """A resolved provision + known lifecycle evt → <textualMod> inside <passiveModifications>."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data

    events = [AmendmentEvent(provision="s 6", effect="am", act_number=99, act_year=2010)]
    xml = _build_tree_with_section(meta, "6")

    # Inject lifecycle so evt_map can resolve the event
    inject_lifecycle(xml, meta, events)
    inject_temporal_data(xml, events)
    inject_passive_mods(xml, events)

    ns = {"akn": AKN_NS}
    pm = xml.find(".//akn:meta/akn:analysis/akn:passiveModifications", ns)
    assert pm is not None, "<passiveModifications> not found"
    mods = pm.findall("akn:textualMod", ns)
    assert len(mods) == 1
    mod = mods[0]
    assert mod.get("eId") == "mod-1"
    src = mod.find("akn:source", ns)
    assert src is not None and src.get("href") == "#evt-amd-1"
    dest = mod.find("akn:destination", ns)
    assert dest is not None and dest.get("href") == "#sec-6"


def test_passive_mod_type_mapping(meta):
    """effect='am' → type='substitution', 'ad' → 'insertion', 'rep' → 'repeal'."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data

    events = [
        AmendmentEvent(provision="s 6", effect="am", act_number=10, act_year=2010),
        AmendmentEvent(provision="s 6", effect="ad", act_number=20, act_year=2011),
        AmendmentEvent(provision="s 6", effect="rep", act_number=30, act_year=2012),
    ]
    xml = _build_tree_with_section(meta, "6")
    inject_lifecycle(xml, meta, events)
    inject_temporal_data(xml, events)
    inject_passive_mods(xml, events)

    ns = {"akn": AKN_NS}
    mods = xml.findall(".//akn:passiveModifications/akn:textualMod", ns)
    assert len(mods) == 3
    types = [m.get("type") for m in mods]
    assert "substitution" in types
    assert "insertion" in types
    assert "repeal" in types


def test_passive_mod_unresolved_skip(meta):
    """Unknown provision (no matching eId) → no <textualMod>, mods_unresolved incremented."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data
    from lexau.models import ParseReport

    # Only sec-6 exists in the tree; provision "s 99" won't resolve
    events = [AmendmentEvent(provision="s 99", effect="am", act_number=99, act_year=2010)]
    xml = _build_tree_with_section(meta, "6")
    inject_lifecycle(xml, meta, events)
    inject_temporal_data(xml, events)

    report = ParseReport(act_name="Test")
    inject_passive_mods(xml, events, report=report)

    ns = {"akn": AKN_NS}
    # <analysis> must not be emitted (zero resolved)
    analysis = xml.find(".//akn:meta/akn:analysis", ns)
    assert analysis is None
    assert report.mods_unresolved == 1
    assert report.mods_resolved == 0


def test_passive_mod_not_applied(meta):
    """applied=False → event is skipped, no <textualMod> emitted."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data
    from lexau.models import ParseReport

    events = [AmendmentEvent(provision="s 6", effect="am", act_number=99, act_year=2010, applied=False)]
    xml = _build_tree_with_section(meta, "6")
    inject_lifecycle(xml, meta, events)
    inject_temporal_data(xml, events)

    report = ParseReport(act_name="Test")
    inject_passive_mods(xml, events, report=report)

    ns = {"akn": AKN_NS}
    analysis = xml.find(".//akn:meta/akn:analysis", ns)
    assert analysis is None
    assert report.mods_resolved == 0


def test_passive_mod_empty_omitted(meta):
    """Zero resolved events → <analysis> element is NOT inserted into <meta>."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data

    # No events at all
    xml = _build_tree_with_section(meta, "6")
    inject_passive_mods(xml, [])

    ns = {"akn": AKN_NS}
    analysis = xml.find(".//akn:meta/akn:analysis", ns)
    assert analysis is None


# --- Task 11: <meta> element order per AKN 3.0 ---

def test_meta_element_order_analysis_before_temporal_data(meta):
    """<analysis> must precede <temporalData> in <meta> per AKN 3.0 XSD."""
    from lexau.builder import inject_passive_mods, inject_lifecycle, inject_temporal_data

    events = [AmendmentEvent(provision="s 6", effect="am", act_number=99, act_year=2010)]
    xml = _build_tree_with_section(meta, "6")
    inject_lifecycle(xml, meta, events)
    inject_temporal_data(xml, events)
    inject_passive_mods(xml, events)

    ns = {"akn": AKN_NS}
    meta_el = xml.find(".//akn:meta", ns)
    children = list(meta_el)

    # Find indices of analysis and temporalData
    analysis_idx = None
    temporal_idx = None
    for i, child in enumerate(children):
        if child.tag == f"{{{AKN_NS}}}analysis":
            analysis_idx = i
        elif child.tag == f"{{{AKN_NS}}}temporalData":
            temporal_idx = i

    # Both must exist
    assert analysis_idx is not None, "<analysis> not found in <meta>"
    assert temporal_idx is not None, "<temporalData> not found in <meta>"

    # <analysis> must come before <temporalData> (per AKN 3.0 XSD)
    assert analysis_idx < temporal_idx, (
        f"<analysis> (index {analysis_idx}) must precede <temporalData> "
        f"(index {temporal_idx}) per AKN 3.0 XSD"
    )


# ---------------------------------------------------------------------------
# Task 8: <quotedStructure> detection
# ---------------------------------------------------------------------------

def test_quoted_structure_detected(meta):
    """BODY('"') SECTION BODY('"') → <quotedStructure> wrapping <section>."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="5", heading="Amendments"),
        ParsedParagraph(ElementType.BODY, text='"'),
        ParsedParagraph(ElementType.SECTION, number="3A", heading="New provision"),
        ParsedParagraph(ElementType.BODY, text='"'),
    ]
    b = AknBuilder(meta)
    for p in paragraphs:
        b.add(p)
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    qs = xml.find(".//akn:quotedStructure", ns)
    assert qs is not None, "<quotedStructure> not found"
    inner_section = qs.find("akn:section", ns)
    assert inner_section is not None, "<section> not found inside <quotedStructure>"


def test_quoted_structure_content(meta):
    """Inner content of <quotedStructure> is valid AKN — has correct attributes."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="10", heading="Modification"),
        ParsedParagraph(ElementType.BODY, text="'"),
        ParsedParagraph(ElementType.SECTION, number="10A", heading="Inserted section"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="Content here."),
        ParsedParagraph(ElementType.BODY, text="'"),
    ]
    b = AknBuilder(meta)
    for p in paragraphs:
        b.add(p)
    xml, _ = b.build()
    ns = {"akn": AKN_NS}
    qs = xml.find(".//akn:quotedStructure", ns)
    assert qs is not None, "<quotedStructure> not found"
    assert qs.get("startQuote") == "“"
    assert qs.get("endQuote") == "”"
    inner_sec = qs.find("akn:section", ns)
    assert inner_sec is not None
    subsec = inner_sec.find("akn:subsection", ns)
    assert subsec is not None
    p_el = subsec.find(".//akn:p", ns)
    assert p_el is not None and p_el.text == "Content here."


def test_quoted_structure_skipped(meta):
    """Multi-provision quote (SECTION SECTION between markers) → quoted_structures_unhandled incremented."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="5", heading="Amendments"),
        ParsedParagraph(ElementType.BODY, text='"'),
        ParsedParagraph(ElementType.SECTION, number="3A", heading="First inserted"),
        ParsedParagraph(ElementType.SECTION, number="3B", heading="Second inserted"),
        ParsedParagraph(ElementType.BODY, text='"'),
    ]
    b = AknBuilder(meta)
    for p in paragraphs:
        b.add(p)
    xml, report = b.build_with_report({})
    assert report.quoted_structures_unhandled >= 1


def test_figure_emitted(meta):
    """A FIGURE paragraph produces <figure><img> in the AKN output."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"),
        ParsedParagraph(ElementType.FIGURE, text=""),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    fig = xml.find(".//akn:figure", ns)
    assert fig is not None
    img = fig.find("akn:img", ns)
    assert img is not None
    src = img.get("src", "")
    assert src.startswith("corpus/images/privacy-act-1988-fig-")
    assert src.endswith(".png")
    assert img.get("alt") == ""


def test_figure_src_increments(meta):
    """Multiple FIGURE paragraphs produce sequentially numbered img src attributes."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"),
        ParsedParagraph(ElementType.FIGURE, text=""),
        ParsedParagraph(ElementType.FIGURE, text=""),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    imgs = xml.findall(".//akn:figure/akn:img", ns)
    assert len(imgs) == 2
    assert imgs[0].get("src") == "corpus/images/privacy-act-1988-fig-1.png"
    assert imgs[1].get("src") == "corpus/images/privacy-act-1988-fig-2.png"


def test_figures_found_in_report(meta):
    """figures_found in ParseReport counts FIGURE paragraphs."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"),
        ParsedParagraph(ElementType.FIGURE, text=""),
        ParsedParagraph(ElementType.FIGURE, text=""),
    ]
    b = AknBuilder(meta)
    for p in paragraphs:
        b.add(p)
    _, report = b.build_with_report({})
    assert report.figures_found == 2


def test_italic_span_emits_i_element(meta):
    """A ParsedParagraph with an italic span produces <p><i>...</i>...</p>."""
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="personal information means something",
        spans=[
            InlineSpan(text="personal information", italic=True),
            InlineSpan(text=" means something"),
        ],
    )
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"),
        p,
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    assert p_el is not None
    i_el = p_el.find(f"{{{AKN_NS}}}i")
    assert i_el is not None
    assert i_el.text == "personal information"
    assert i_el.tail == " means something"


def test_bold_span_emits_b_element(meta):
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="Important note",
        spans=[InlineSpan(text="Important note", bold=True)],
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    assert p_el.find(f"{{{AKN_NS}}}b") is not None
    assert p_el.find(f"{{{AKN_NS}}}b").text == "Important note"


def test_plain_spans_no_children(meta):
    """Unformatted spans produce plain p.text, no children."""
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="This Act is the Privacy Act 1988.",
        spans=[InlineSpan(text="This Act is the Privacy Act 1988.")],
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    assert len(list(p_el)) == 0
    assert p_el.text == "This Act is the Privacy Act 1988."


def test_no_spans_falls_back_to_plain_text(meta):
    """ParsedParagraph with empty spans list behaves exactly as before."""
    p = ParsedParagraph(
        ElementType.BODY,
        text="This Act is the Privacy Act 1988.",
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    assert len(list(p_el)) == 0
    assert p_el.text == "This Act is the Privacy Act 1988."


def test_superscript_span_emits_sup_element(meta):
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="CO2",
        spans=[InlineSpan(text="CO"), InlineSpan(text="2", superscript=True)],
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Formula"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    sup_el = p_el.find(f"{{{AKN_NS}}}sup")
    assert sup_el is not None
    assert sup_el.text == "2"


def test_bold_italic_span_emits_nested_b_i(meta):
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="critical term",
        spans=[InlineSpan(text="critical term", bold=True, italic=True)],
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Definitions"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    b_el = p_el.find(f"{{{AKN_NS}}}b")
    assert b_el is not None
    i_el = b_el.find(f"{{{AKN_NS}}}i")
    assert i_el is not None
    assert i_el.text == "critical term"


def test_adjacent_same_formatting_spans_no_spurious_space_on_pretty_print(meta):
    """Real case (Task 1 triage 2026-09-08, wp_garble Group C, 819 records,
    62% of all wp_garble -- the highest-volume converter defect this task
    fixes): excise-tariff-amendment-act-1990's long title has "Excise Tariff
    Amendment Act 1990" split across two adjacent DOCX <w:r> runs with
    IDENTICAL formatting (both bold), separated only by an intervening
    <w:bookmarkStart>/<w:bookmarkEnd> pair (Word's auto-inserted "_GoBack"
    cursor-position bookmark -- no actual whitespace between the runs).
    _emit_p_inline previously created two SIBLING <b> elements with no tail
    text between them; lxml's pretty_print serializer (used for the
    persisted corpus/xml/*.xml output, src/lexau/corpus.py:77) then inserts
    newline+indent whitespace as that missing tail, corrupting "Excise" into
    "E xcise" once whitespace-normalised. Also confirmed in
    acts-and-instruments-(framework-reform)-act-2015's long title
    ("sunsetting" -> "sunset" + "ting" split across two bold runs).
    """
    from lexau.parser import InlineSpan
    p = ParsedParagraph(
        ElementType.BODY,
        text="Excise Tariff Amendment Act 1990",
        spans=[
            InlineSpan(text="E", bold=True),
            InlineSpan(text="xcise Tariff Amendment Act 1990", bold=True),
        ],
    )
    xml, _ = build_xml(meta, [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        p,
    ])
    ns = {"akn": AKN_NS}
    p_el = xml.find(".//akn:section/akn:content/akn:p", ns)
    b_els = p_el.findall(f"{{{AKN_NS}}}b")
    assert [b.text for b in b_els] == ["E", "xcise Tariff Amendment Act 1990"]
    # The load-bearing assertion: no element between two same-formatting
    # runs may be left with tail=None, since lxml's pretty-print serializer
    # (the one actually used for persisted output) fills a None tail with
    # indentation whitespace, corrupting contiguous source text.
    pretty = etree.tostring(p_el, pretty_print=True)
    reparsed = etree.fromstring(pretty)
    assert "".join(reparsed.itertext()) == "Excise Tariff Amendment Act 1990"


def test_build_with_report_list_defs_found(meta):
    """build_with_report counts list-form definitions in list_defs_found."""
    from unittest.mock import patch, MagicMock
    from lexau.endnote_parser import EndnoteResult
    from lexau.parser import InlineSpan

    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"))
    b.add(ParsedParagraph(ElementType.SUBSECTION, number="1", text=""))
    # List-form definition paragraph
    b.add(ParsedParagraph(ElementType.BODY, text="agency means:"))
    # Followed by paragraph items
    b.add(ParsedParagraph(ElementType.PARAGRAPH, number="a", text="a body corporate; or"))
    b.add(ParsedParagraph(ElementType.PARAGRAPH, number="b", text="a natural person."))

    corpus_index: dict = {}
    with patch("lexau.builder.parse_endnotes", return_value=EndnoteResult([], [])):
        _, report = b.build_with_report(corpus_index, last_volume_path=None)

    assert report.list_defs_found >= 1


def test_build_with_report_list_defs_completed(meta):
    """build_with_report folds orphaned list content into a truncated <def>
    via complete_list_definitions, and reports the count."""
    from unittest.mock import patch
    from lexau.endnote_parser import EndnoteResult

    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"))
    b.add(ParsedParagraph(ElementType.SUBSECTION, number="1", text=""))
    b.add(ParsedParagraph(ElementType.BODY, text="eligible entity means any of the following:"))
    b.add(ParsedParagraph(ElementType.PARAGRAPH, number="a", text="a body corporate; or"))
    b.add(ParsedParagraph(ElementType.PARAGRAPH, number="b", text="a natural person."))

    corpus_index: dict = {}
    with patch("lexau.builder.parse_endnotes", return_value=EndnoteResult([], [])):
        xml, report = b.build_with_report(corpus_index, last_volume_path=None)

    assert report.list_defs_completed >= 1
    ns = {"akn": AKN_NS}
    def_el = xml.find(".//akn:def", ns)
    assert def_el is not None
    text = "".join(def_el.itertext())
    assert "a body corporate" in text
    assert "a natural person" in text


def test_asterisk_ref_injected_end_to_end(meta):
    """build_with_report resolves *term usages against the term registry into <ref> links."""
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Dictionary"))
    b.add(ParsedParagraph(ElementType.BODY, text='"entity" means a person or organisation.'))
    b.add(ParsedParagraph(ElementType.SECTION, number="2", heading="Obligations"))
    b.add(ParsedParagraph(ElementType.BODY, text="The *entity must comply with this Act."))
    corpus_index: dict = {}
    xml, report = b.build_with_report(corpus_index)

    assert report.asterisk_resolved == 1
    assert report.asterisk_unresolved == 0
    refs = xml.findall(f".//{{{AKN_NS}}}ref")
    assert any(r.get("href") == "#term-entity" for r in refs)


def test_asterisk_ref_and_narrative_guard_combined_end_to_end(meta):
    """Both features exercised in one builder run within one definitions section:
    a genuine quoted definition resolves an *entity usage into a <ref>, while a
    narrative-prose candidate in the same section (an embedded relative clause,
    "(who may include ...)") is correctly rejected by the narrative guard and
    injects no spurious <term>."""
    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Dictionary"))
    b.add(ParsedParagraph(ElementType.BODY, text='"entity" means a person or organisation.'))
    b.add(ParsedParagraph(
        ElementType.BODY,
        text="a person (who may include the trustee) is empowered to exercise any power of appointment.",
    ))
    b.add(ParsedParagraph(ElementType.SECTION, number="2", heading="Obligations"))
    b.add(ParsedParagraph(ElementType.BODY, text="The *entity must comply with this Act."))
    corpus_index: dict = {}
    xml, report = b.build_with_report(corpus_index)

    # Genuine usage: asterisk-ref resolved against the registry.
    assert report.asterisk_resolved == 1
    assert report.asterisk_unresolved == 0
    refs = xml.findall(f".//{{{AKN_NS}}}ref")
    assert any(r.get("href") == "#term-entity" for r in refs)

    # Rejected narrative candidate: no spurious second <term> injected.
    assert report.terms_found == 1
    term_els = xml.findall(f".//{{{AKN_NS}}}term")
    assert len(term_els) == 1
    assert term_els[0].get("refersTo") == "#term-entity"
    assert not any("who may" in (t.text or "") for t in term_els)


def test_build_with_report_inline_formatted_count(meta):
    """build_with_report counts <p> elements with inline markup in inline_formatted."""
    from lexau.parser import InlineSpan

    b = AknBuilder(meta)
    b.add(ParsedParagraph(ElementType.SECTION, number="6", heading="Definitions"))
    b.add(ParsedParagraph(
        ElementType.BODY,
        text="personal information means something",
        spans=[
            InlineSpan(text="personal information", italic=True),
            InlineSpan(text=" means something"),
        ],
    ))
    corpus_index: dict = {}
    with patch("lexau.builder.parse_endnotes", return_value=EndnoteResult([], [])):
        _, report = b.build_with_report(corpus_index, last_volume_path=None)

    assert report.inline_formatted >= 1


def test_list_item_inline_formatting_not_emitted_known_limitation(meta):
    """v0.6.0: LIST_ITEM <p> body is plain text even when spans carry formatting.
    _emit_p_inline is not called in the LIST_ITEM branch of build() — known limitation."""
    from lexau.parser import InlineSpan
    from lexau.builder import _emit_p_inline
    # _emit_p_inline itself works fine for any paragraph type
    p_with_bold = ParsedParagraph(
        ElementType.LIST_ITEM,
        number="1",
        text="bold text",
        spans=[InlineSpan(text="bold text", bold=True)],
    )
    p_el = etree.Element(f"{{{AKN_NS}}}p")
    _emit_p_inline(p_el, p_with_bold)
    # _emit_p_inline correctly emits <b> — but the LIST_ITEM branch in build() never calls this
    assert p_el.find(f"{{{AKN_NS}}}b") is not None
    # The actual limitation: build() assigns p_el.text directly for LIST_ITEM
    p_el_direct = etree.Element(f"{{{AKN_NS}}}p")
    p_el_direct.text = p_with_bold.text  # what build() actually does
    assert p_el_direct.find(f"{{{AKN_NS}}}b") is None, "LIST_ITEM branch uses .text, not _emit_p_inline"


def test_bold_superscript_span_emits_b_only_known_limitation(meta):
    """v0.6.0: bold+superscript span emits <b> only — superscript is dropped.
    Only bold+italic nesting is handled; other combinations are not spec'd."""
    from lexau.parser import InlineSpan
    from lexau.builder import _emit_p_inline
    p = ParsedParagraph(
        ElementType.SECTION,
        number="1",
        text="CO2",
        spans=[InlineSpan(text="CO2", bold=True, superscript=True)],
    )
    p_el = etree.Element(f"{{{AKN_NS}}}p")
    _emit_p_inline(p_el, p)
    b_el = p_el.find(f"{{{AKN_NS}}}b")
    assert b_el is not None, "bold wins over superscript"
    assert b_el.text == "CO2"
    # superscript is dropped — known v0.6.0 limitation (plan only spec'd bold+italic nesting)
    assert p_el.find(f"{{{AKN_NS}}}sup") is None, "superscript not emitted when bold also set (v0.6.0 limitation)"


# --- v0.9.1: real figure src + dimensions from captured blobs -----------------

_FIG_FIXTURES = Path(__file__).parent / "fixtures" / "figures"


def _figure_meta(name: str) -> ActMetadata:
    # ActMetadata.safe_name = name.lower with spaces/slashes -> '-', so a name
    # with no spaces/slashes passes through verbatim.
    return ActMetadata(
        name=name,
        title_id="C0000A00000",
        comp_id="C0000C00000",
        comp_num="1",
        year=2000,
        number=1,
        effective_date=_date(2020, 1, 1),
    )


def _convert_figure_volumes(docx_paths, meta, images_out):
    """Feed one or more figure fixture DOCX volumes through the real
    docx_reader -> AknBuilder path and return (report, xml_string)."""
    from dataclasses import replace as _replace
    from docx import Document as _Document
    from lexau.docx_reader import iter_paragraphs as _iter

    b = AknBuilder(meta, images_out=images_out)
    # A leading SECTION so the figure lands in <body>, not <preface>
    # (_split_stream routes everything before the first structural element to
    # the preface, where _figures_found never advances).
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"))
    for vol_idx, path in enumerate(docx_paths):
        for p in _iter(_Document(str(path))):
            b.add(_replace(p, volume_index=vol_idx))
    root, report = b.build_with_report({})
    return report, etree.tostring(root, encoding="unicode")


def test_builder_writes_real_figure_src(tmp_path):
    meta = _figure_meta("demo_act")
    report, xml = _convert_figure_volumes(
        [str(_FIG_FIXTURES / "one_png.docx")], meta, tmp_path
    )
    assert 'src="corpus/images/demo_act-fig-1.png"' in xml
    assert "width=" in xml and "height=" in xml
    assert report.figures_found == 1
    assert report.figures_raster == 1
    assert report.figures_converted == 0
    assert report.figures_placeholder == 0
    assert (tmp_path / "demo_act-fig-1.png").exists()


def test_builder_figure_src_verbatim_not_forced_png(tmp_path):
    """The <img src> is whatever FigureResult.src says. A .gif raster keeps
    its .gif extension; the builder must not hardcode .png."""
    from unittest.mock import patch
    from lexau.figures import FigureResult

    meta = _figure_meta("gif_act")
    fake = [[FigureResult("corpus/images/gif_act-fig-1.gif", "raster", 12, 34)]]
    with patch("lexau.builder.materialise_figures", return_value=fake) as m:
        report, xml = _convert_figure_volumes(
            [str(_FIG_FIXTURES / "one_png.docx")], meta, tmp_path
        )
    assert m.call_count == 1
    assert 'src="corpus/images/gif_act-fig-1.gif"' in xml
    assert 'width="12"' in xml and 'height="34"' in xml
    assert report.figures_raster == 1


def test_builder_materialises_figures_once_across_volumes(tmp_path, monkeypatch):
    """ITAA-style multi-volume: materialise_figures is called ONCE over the
    concatenated blob list in global order, so volume 2's image is fig-2 and
    volume 1's fig-1 is not overwritten by a per-volume restart."""
    import lexau.builder as builder_mod

    calls: list = []
    real = builder_mod.materialise_figures

    def spy(slug, safe_name, figures, out_dir):
        calls.append([list(row) for row in figures])
        return real(slug, safe_name, figures, out_dir)

    monkeypatch.setattr(builder_mod, "materialise_figures", spy)

    meta = _figure_meta("two_vol_act")
    report, xml = _convert_figure_volumes(
        [
            str(_FIG_FIXTURES / "one_png.docx"),
            str(_FIG_FIXTURES / "one_png.docx"),
        ],
        meta,
        tmp_path,
    )
    assert len(calls) == 1, "one materialise call per Act, not per volume"
    assert len(calls[0]) == 2, "both volumes' figures in one concatenated list"
    assert report.figures_found == 2
    assert report.figures_raster == 2
    assert 'src="corpus/images/two_vol_act-fig-1.png"' in xml
    assert 'src="corpus/images/two_vol_act-fig-2.png"' in xml
    assert (tmp_path / "two_vol_act-fig-1.png").exists()
    assert (tmp_path / "two_vol_act-fig-2.png").exists()


def test_builder_empty_blob_figure_is_placeholder(tmp_path):
    """A FIGURE paragraph with no captured image keeps the computed
    placeholder src and counts as figures_placeholder."""
    meta = _figure_meta("ph_act")
    b = AknBuilder(meta, images_out=tmp_path)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"))
    b.add(ParsedParagraph(ElementType.FIGURE, text=""))
    root, report = b.build_with_report({})
    xml = etree.tostring(root, encoding="unicode")
    assert 'src="corpus/images/ph_act-fig-1.png"' in xml
    assert report.figures_found == 1
    assert report.figures_placeholder == 1
    assert report.figures_raster == 0
    assert not (tmp_path / "ph_act-fig-1.png").exists()


def test_figure_paragraph_with_two_blips_captures_only_first(tmp_path):
    """A FIGURE <w:p> that carries two inline images (the
    excise-tariff-act-1921 / corporate-law-economic-reform-program-act-1999
    pattern, where the 2nd blob is a byte-identical preview of the next
    figure) must yield exactly one blob per FIGURE paragraph, so no orphan
    ``-fig-1b`` file is written."""
    import io
    from docx import Document as _Document
    from lexau.docx_reader import iter_paragraphs as _iter

    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d494844520000000100000001080600000"
        "01f15c4890000000a49444154789c6300010000050001"
        "0d0a2db40000000049454e44ae426082"
    )
    src = tmp_path / "two_blip.docx"
    doc = _Document()
    doc.add_paragraph("Intro prose so the doc has a non-figure paragraph.")
    para = doc.add_paragraph()
    para.add_run().add_picture(io.BytesIO(png))
    para.add_run().add_picture(io.BytesIO(png))
    doc.save(src)

    figs = [
        p for p in _iter(_Document(str(src)))
        if p.element_type == ElementType.FIGURE
    ]
    assert len(figs) == 1
    assert len(figs[0].image_blobs) == 1

    meta = _figure_meta("two_blip_act")
    images_out = tmp_path / "images"
    b = AknBuilder(meta, images_out=images_out)
    b.add(ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"))
    for p in _iter(_Document(str(src))):
        b.add(p)
    root, report = b.build_with_report({})

    assert report.figures_found == 1
    assert (images_out / "two_blip_act-fig-1.png").exists()
    assert not (images_out / "two_blip_act-fig-1b.png").exists()


def test_schedule_date_line_not_clause_heading(meta):
    # Date-like lines starting with number+month should not be fabricated as clause headings.
    # Regression test for lex-au-explorer FASA definitions ("30 June 2000 rate, in relation to...").
    # Positive control: "30 Standard rate" should still become a clause heading.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Definitions", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="30 Standard rate"),
        ParsedParagraph(ElementType.BODY, text="30 June 2000 rate, in relation to an individual, means..."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    # "30 Standard rate" should still become a clause heading
    clause = xml.find(".//akn:hcontainer[@name='clause'][@eId='schedule-1__clause-30']", ns)
    assert clause is not None, "Positive control: '30 Standard rate' should fabricate clause-30"

    # The clause heading should contain "Standard rate", not "June 2000"
    heading = clause.find("akn:heading", ns)
    assert heading is not None
    assert "Standard rate" in heading.text

    # The date line should appear as plain prose in the schedule content
    prose_paras = xml.findall(".//akn:hcontainer[@name='schedule']/akn:content/akn:p", ns)
    date_prose = [p for p in prose_paras if "June 2000" in (p.text or "")]
    assert len(date_prose) > 0, "Date line should appear as prose in schedule content"


# ---------------------------------------------------------------------------
# Task 6 / §3 B4 — schedule quoted-structure + amendment-instruction awareness
#
# Four collision mechanisms from the P1 note
# (docs/superpowers/notes/2026-09-07-p1-schedule-structure.md):
#   M3b  — ItemHead instruction number vs a quoted provision's own section number
#   M3bt — an embedded schedule TOC line vs the real section it indexes
#   M2   — schedule Part boundary resets the item numbering, no grouping kept
#   M1   — amended-Act citation heading resets the item numbering
# Shapes below are taken verbatim from the real DOCX paragraph streams of the
# Acts the brief names (verified via docx_reader against corpus/docx).
# ---------------------------------------------------------------------------


def _all_eids(xml):
    return [el.get("eId") for el in xml.iter() if el.get("eId")]


def test_m3b_item_head_and_quoted_section_do_not_collide(meta):
    # agricultural-and-veterinary-chemicals-legislation-amendment-act-2013,
    # schedule 1. Real stream: ItemHead "9  Subsection 3(1) ..." earlier, and
    # later ItemHead "30  Section 9 ..." + Item "Repeal the section, substitute:"
    # + ActHead 5 SECTION 9 "Explanation of Part" (quoted replacement law).
    # Today both write schedule-1__clause-9.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Approvals", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Agricultural and Veterinary Chemicals Code Act 1994", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="9  Subsection\xa03(1) of the Code set out in the Schedule (definition of established standard)", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Repeal the definition, substitute:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="established standard has the meaning given by subsection\xa08U(7).", raw_style="Definition"),
        ParsedParagraph(ElementType.BODY, text="30  Section\xa09 of the Code set out in the Schedule", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Repeal the section, substitute:", raw_style="Item"),
        ParsedParagraph(ElementType.SECTION, number="9", heading="Explanation of Part", raw_style="ActHead 5"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="This Part contains provisions relating to:", raw_style="subsection"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    # The two ItemHead lines are amendment items, not schedule clauses.
    items = xml.findall(".//akn:hcontainer[@name='item']", ns)
    assert len(items) == 2, f"expected 2 <hcontainer name='item'>, got {len(items)}"
    assert items[0].find("akn:num", ns).text == "9"
    assert items[1].find("akn:num", ns).text == "30"
    # Keyed on a monotonic index, not the literal instruction number (P1 B4 item 2).
    assert items[0].get("eId").endswith("__item-1")
    assert items[1].get("eId").endswith("__item-2")

    # No fabricated schedule clause anywhere.
    assert xml.findall(".//akn:hcontainer[@name='clause']", ns) == []

    # The quoted replacement provision keeps its own hierarchy inside a
    # <quotedStructure>, scoped under item 2 — never schedule-1__clause-9.
    qs = xml.findall(".//akn:quotedStructure", ns)
    assert len(qs) == 2, f"expected 2 <quotedStructure>, got {len(qs)}"
    inner_section = qs[1].find("akn:section", ns)
    assert inner_section is not None, "<section> not found inside <quotedStructure>"
    assert inner_section.find("akn:num", ns).text == "9"
    assert inner_section.get("eId").startswith(items[1].get("eId") + "__")
    assert "clause-9" not in inner_section.get("eId")
    assert xml.find(".//akn:hcontainer[@eId='schedule-1__clause-9']", ns) is None

    # Word-for-word: every source line still present.
    text = " ".join(" ".join(xml.itertext()).split())
    assert "Explanation of Part" in text
    assert "established standard has the meaning given by subsection 8U(7)." in text
    assert "This Part contains provisions relating to:" in text


def test_m3bt_schedule_toc_lines_are_skipped(meta):
    # competition-and-consumer-act-2010 schedule 2 (the Australian Consumer Law):
    # a rendered schedule TOC (Special TOC 1/2/5) precedes the real Chapter /
    # Part / section headings. Today every TOC line fabricates a clause that
    # collides with the section it indexes.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—The Australian Consumer Law", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Chapter\xa01—Introduction", raw_style="Special TOC 1"),
        ParsedParagraph(ElementType.BODY, text="1\tApplication of this Schedule", raw_style="Special TOC 5"),
        ParsedParagraph(ElementType.BODY, text="2\tDefinitions", raw_style="Special TOC 5"),
        ParsedParagraph(ElementType.CHAPTER, number="1", heading="Introduction", raw_style="ActHead 2"),
        ParsedParagraph(ElementType.SECTION, number="1", heading="Application of this Schedule", raw_style="ActHead 5"),
        ParsedParagraph(ElementType.BODY, text="This Schedule applies to the extent provided by:", raw_style="subsection"),
        ParsedParagraph(ElementType.SECTION, number="2", heading="Definitions", raw_style="ActHead 5"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    # The TOC lines produce no element at all.
    sched = xml.find(".//akn:hcontainer[@name='schedule']", ns)
    text = " ".join(" ".join(sched.itertext()).split())
    assert "1\tApplication of this Schedule" not in "".join(sched.itertext())
    assert text.count("Application of this Schedule") == 1, text
    assert text.count("Definitions") == 1, text

    clauses = sched.findall(".//akn:hcontainer[@name='clause']", ns)
    assert len(clauses) == 2, f"expected exactly 2 clauses (sections 1 and 2), got {len(clauses)}"
    assert [c.find("akn:num", ns).text for c in clauses] == ["1", "2"]


def test_m2_part_boundary_gives_repeated_item_numbers_distinct_eids(meta):
    # fair-work-amendment-(protecting-vulnerable-workers)-act-2017 schedule 1:
    # Part 1 instruction items run 1..13, Part 8 restarts at 1. Today both
    # write schedule-1__clause-1.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.PART, number="1", heading="Increasing maximum penalties", raw_style="ActHead 7"),
        ParsedParagraph(ElementType.BODY, text="Fair Work Act 2009", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  Section\xa012", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="serious contravention has the meaning given by section\xa0557A.", raw_style="Definition"),
        ParsedParagraph(ElementType.PART, number="8", heading="Records", raw_style="ActHead 7"),
        ParsedParagraph(ElementType.BODY, text="Fair Work Act 2009", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  At the end of subsection\xa0535(3)", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Add:", raw_style="Item"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    # Both Part headings survive as grouping wrappers.
    parts = xml.findall(".//akn:hcontainer[@name='schedule']//akn:hcontainer[@name='part']", ns)
    assert len(parts) == 2, f"expected 2 schedule part wrappers, got {len(parts)}"
    assert parts[0].get("eId") == "schedule-1__part-1"
    assert parts[1].get("eId") == "schedule-1__part-8"

    items = xml.findall(".//akn:hcontainer[@name='item']", ns)
    assert len(items) == 2
    a, b = items[0].get("eId"), items[1].get("eId")
    assert a != b, f"item-1 of Part 1 and item-1 of Part 8 still collide on {a}"
    assert a.startswith("schedule-1__part-1__"), a
    assert b.startswith("schedule-1__part-8__"), b
    # The literal instruction number is preserved in <num> either way.
    assert [i.find("akn:num", ns).text for i in items] == ["1", "1"]


def test_m1_amended_act_boundary_gives_repeated_item_numbers_distinct_eids(meta):
    # statute-law-revision-act-2012 schedule 2: a flat list of typographical
    # corrections grouped under ActHead 9 amended-Act name headings, the item
    # count restarting per Act. Today both write schedule-1__clause-1 and the
    # Act-name line is buried as a <p> inside the first one's <content>.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Amendment of amending Acts", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Electoral and Referendum Amendment (Enrolment) Act 2011", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  After subparagraph\xa0110(4)(b)(iv)", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Insert “and”.", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="Fair Work (State Referral) Act 2009", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  Item\xa0215 of Schedule\xa01", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Omit “DPP”, substitute “Director”.", raw_style="Item"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    acts = xml.findall(".//akn:hcontainer[@name='amendedAct']", ns)
    assert len(acts) == 2, f"expected 2 amended-Act wrappers, got {len(acts)}"
    assert acts[0].get("eId") == "schedule-1__amdact-1"
    assert acts[1].get("eId") == "schedule-1__amdact-2"
    # The Act-name line becomes the wrapper's heading, not buried prose.
    assert acts[0].find("akn:heading", ns).text == "Electoral and Referendum Amendment (Enrolment) Act 2011"

    items = xml.findall(".//akn:hcontainer[@name='item']", ns)
    assert len(items) == 2
    a, b = items[0].get("eId"), items[1].get("eId")
    assert a != b, f"item 1 of each amended Act still collide on {a}"
    assert a.startswith("schedule-1__amdact-1__"), a
    assert b.startswith("schedule-1__amdact-2__"), b


_PLAIN_SCHEDULE_BASELINE = """<attachments xmlns="http://docs.oasis-open.org/legaldocml/ns/akn/3.0">
  <attachment>
    <hcontainer name="schedule" eId="schedule-1">
      <num>1</num>
      <heading>Essential principles</heading>
      <hcontainer name="clause" eId="schedule-1__clause-1">
        <num>1</num>
        <heading>General requirements</heading>
        <content>
          <p>A device must be designed to be safe.</p>
        </content>
      </hcontainer>
      <hcontainer name="clause" eId="schedule-1__clause-7">
        <num>7</num>
        <heading>Chemical properties</heading>
        <hcontainer name="subclause" eId="schedule-1__clause-7__subclause-7-1">
          <num>7.1</num>
          <heading>Choice of materials</heading>
          <content>
            <p>Materials must be biocompatible.</p>
          </content>
          <paragraph eId="schedule-1__clause-7__subclause-7-1__para-a">
            <num>a</num>
            <content>
              <p>toxicity of materials</p>
            </content>
            <subparagraph eId="schedule-1__clause-7__subclause-7-1__para-a__subpara-i">
              <num>i</num>
              <content>
                <p>acute toxicity</p>
              </content>
            </subparagraph>
          </paragraph>
          <hcontainer name="subclause" eId="schedule-1__clause-7__subclause-7-1__subclause-2">
            <num>2</num>
            <content>
              <p>The manufacturer must document this.</p>
            </content>
          </hcontainer>
          <table>
            <tr>
              <td>Item</td>
              <td>Requirement</td>
            </tr>
            <tr>
              <td>1</td>
              <td>Sterility</td>
            </tr>
          </table>
          <content>
            <p>Trailing prose after the table.</p>
          </content>
        </hcontainer>
        <content>
          <p>Note: see clause 8.</p>
        </content>
      </hcontainer>
    </hcontainer>
  </attachment>
</attachments>
"""


def test_b4_leaves_plain_non_amending_schedule_byte_identical(meta):
    # Regression gate: a schedule with no amendment instructions, no quoted
    # structure and no TOC must serialise exactly as it did at b88624a.
    # Baseline string captured from the pre-B4 builder.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Essential principles", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="1  General requirements"),
        ParsedParagraph(ElementType.BODY, text="A device must be designed to be safe."),
        ParsedParagraph(ElementType.SECTION, number="7", heading="Chemical properties"),
        ParsedParagraph(ElementType.SECTION, number="7.1", heading="Choice of materials"),
        ParsedParagraph(ElementType.BODY, text="Materials must be biocompatible."),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="toxicity of materials"),
        ParsedParagraph(ElementType.SUBPARAGRAPH, number="i", text="acute toxicity"),
        ParsedParagraph(ElementType.SUBSECTION, number="2", text="The manufacturer must document this."),
        ParsedParagraph(ElementType.TABLE, table_rows=[["Item", "Requirement"], ["1", "Sterility"]]),
        ParsedParagraph(ElementType.BODY, text="Trailing prose after the table."),
        ParsedParagraph(ElementType.NOTE, text="Note: see clause 8."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    att = xml.find(f".//{{{AKN_NS}}}attachments")
    assert etree.tostring(att, pretty_print=True).decode() == _PLAIN_SCHEDULE_BASELINE


# ---------------------------------------------------------------------------
# Task 6 fix round — quoted-run termination (reviewer Finding 1)
#
# `_ends_quoted_run` closed a quoted run on a structural break or a margin note
# but not on a FRESH amendment instruction, so a second instruction inside one
# item was swallowed into the first <quotedStructure> and rendered as quoted
# law. The inverse also held: any `Item`-styled colon-terminated line could open
# a <quotedStructure> around ordinary transitional prose.
#
# Shapes below come from the real DOCX paragraph streams of the Acts named in
# the review, read via docx_reader against corpus/docx.
# ---------------------------------------------------------------------------


def test_two_instructions_in_one_item_open_separate_quoted_structures(meta):
    # parliamentary-business-resources-(consequential-and-transitional-provisions)-act-2017
    # schedule 1 item 1, and the same shape in NDIS (worker-screening-database)
    # 2019 schedule 1 item 1: one ItemHead, an "Omit:" block of old text, then a
    # bare "substitute:" introducing the replacement block.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Parliamentary Business Resources Act 2017", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  Section\xa03", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Omit:", raw_style="Item"),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="audits relating to work expenses.", raw_style="SO Para"),
        ParsedParagraph(ElementType.BODY, text="substitute:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="•\tThe Authority has functions relating to:", raw_style="SO Bullet"),
        ParsedParagraph(ElementType.PARAGRAPH, number="b", text="the work resources and travel resources.", raw_style="SO Para"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    items = xml.findall(".//akn:hcontainer[@name='item']", ns)
    assert len(items) == 1, f"expected 1 item, got {len(items)}"
    item = items[0]

    qs = item.findall("akn:quotedStructure", ns)
    assert len(qs) == 2, f"expected 2 <quotedStructure> in the item, got {len(qs)}"
    assert qs[0].get("eId").endswith("__qstr-1")
    assert qs[1].get("eId").endswith("__qstr-2")

    # The instruction lines are the item's own prose, never quoted law.
    for q in qs:
        quoted_text = " ".join(" ".join(q.itertext()).split())
        assert "substitute:" not in quoted_text, f"instruction swallowed into quoted law: {quoted_text!r}"
        assert "Omit:" not in quoted_text
    item_prose = [
        " ".join((p.text or "").split())
        for p in item.findall("akn:content/akn:p", ns)
    ]
    assert "Omit:" in item_prose
    assert "substitute:" in item_prose

    # Old text in the first quoted structure, replacement in the second.
    assert "audits relating to work expenses." in "".join(qs[0].itertext())
    assert "The Authority has functions relating to:" in "".join(qs[1].itertext())
    assert "the work resources and travel resources." in "".join(qs[1].itertext())

    # Word-for-word: nothing lost.
    text = " ".join(" ".join(xml.itertext()).split())
    for fragment in (
        "Omit:", "substitute:",
        "audits relating to work expenses.",
        "The Authority has functions relating to:",
        "the work resources and travel resources.",
    ):
        assert fragment in text


def test_colon_prose_inside_quoted_law_does_not_truncate_the_run(meta):
    # broadcasting-legislation-amendment-(broadcasting-reform)-act-2017: an
    # inserted transitional provision whose own definition intros are Item-styled
    # and end in a colon ("eligible financial year means:", "where:"). These are
    # quoted law, NOT fresh instructions -- the run must run past them.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="1  In the appropriate position", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.SECTION, number="14", heading="Rebate", raw_style="ActHead 5"),
        ParsedParagraph(ElementType.BODY, text="eligible financial year means:", raw_style="Item"),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="the 2013-2014 financial year; or", raw_style="paragraph"),
        ParsedParagraph(ElementType.BODY, text="where:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="days in non-licence period means the number of days in the period:", raw_style="Item"),
        ParsedParagraph(ElementType.PARAGRAPH, number="b", text="beginning on the designated day.", raw_style="paragraph"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    qs = xml.findall(".//akn:quotedStructure", ns)
    assert len(qs) == 1, f"genuine quoted span was split into {len(qs)} -- expected 1"

    quoted = " ".join(" ".join(qs[0].itertext()).split())
    for fragment in (
        "Rebate",
        "eligible financial year means:",
        "the 2013-2014 financial year; or",
        "where:",
        "days in non-licence period means the number of days in the period:",
        "beginning on the designated day.",
    ):
        assert fragment in quoted, f"{fragment!r} fell out of the quoted structure"


def test_item_styled_transitional_prose_does_not_open_a_quoted_structure(meta):
    # The inverse false positive: `Item`-styled application/transitional prose
    # that ends in a colon ("The repeal of ... does not affect:", "In this Part:")
    # is not an amendment instruction and must not wrap the following prose in a
    # <quotedStructure>.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Transitional", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="5  Saving provision", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="The repeal of section\xa010 by this Schedule does not affect:", raw_style="Item"),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="a determination in force immediately before the repeal; or", raw_style="paragraph"),
        ParsedParagraph(ElementType.PARAGRAPH, number="b", text="anything done under such a determination.", raw_style="paragraph"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    assert xml.findall(".//akn:quotedStructure", ns) == [], (
        "transitional prose ending in a colon must not open a <quotedStructure>"
    )
    item = xml.find(".//akn:hcontainer[@name='item']", ns)
    assert item is not None
    text = " ".join(" ".join(item.itertext()).split())
    for fragment in (
        "The repeal of section\xa010 by this Schedule does not affect:".replace("\xa0", " "),
        "a determination in force immediately before the repeal; or",
        "anything done under such a determination.",
    ):
        assert fragment in text


def test_nested_amendment_item_keeps_its_own_instructions_quoted(meta):
    # statute-law-revision genre: an item repeals an amending Act's item and
    # substitutes a replacement, whose own `Special ih` sub-items carry their own
    # `Item` instructions. Those are quoted law -- an "Insert:" after a
    # `Special ih` must NOT close the outer quoted run.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa02—Amendment of amending Acts", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="4  Item\xa0213 of Schedule\xa02", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Repeal the item, substitute:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="213  Subsection\xa045(1A)", raw_style="Special ih"),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="authorised officer means an officer of Customs.", raw_style="Definition"),
        ParsedParagraph(ElementType.BODY, text="213A  Paragraph\xa045(1A)(a)", raw_style="Special ih"),
        ParsedParagraph(ElementType.BODY, text="Omit “Director”, substitute “authority”.", raw_style="Item"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    qs = xml.findall(".//akn:quotedStructure", ns)
    assert len(qs) == 1, f"nested replacement item was split into {len(qs)} quoted structures"
    quoted = " ".join(" ".join(qs[0].itertext()).split())
    for fragment in (
        "213 Subsection 45(1A)",
        "Insert:",
        "authorised officer means an officer of Customs.",
        "213A Paragraph 45(1A)(a)",
        "Omit “Director”, substitute “authority”.",
    ):
        assert fragment in quoted, f"{fragment!r} fell out of the nested quoted item"


# --- v0.3.1 §2: FIGURE paragraphs keep their own provision text ---------------


def _mixed_vml_docx(tmp_path, text, name="mixed_vml.docx"):
    """A one-paragraph DOCX whose single <w:p> carries BOTH ``text`` and a VML
    (<w:pict>/<v:imagedata>) image — the pre-2010 drafting shape behind the 7
    structural-loss Acts in the P3 note."""
    from docx import Document as _Document
    from docx.oxml import parse_xml as _parse_xml
    from docx.oxml.ns import qn as _qn

    from docx.enum.style import WD_STYLE_TYPE

    doc = _Document(str(_FIG_FIXTURES / "one_wmf.docx"))
    # An ActHead-styled section heading ahead of the figure paragraph: it makes
    # the document non-legacy (so parse_paragraph classifies, as in all 7 of
    # the real structural-loss Acts) and opens a section for the eId prefix.
    doc.styles.add_style("ActHead 5", WD_STYLE_TYPE.PARAGRAPH)
    para = doc.paragraphs[0]
    heading = doc.add_paragraph("4A  Rate of pension")
    heading.style = doc.styles["ActHead 5"]
    para._element.addprevious(heading._element)
    for run in list(para._element.findall(_qn("w:r"))):
        para._element.remove(run)
    para.add_run(text)
    para._element.append(
        _parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<w:pict><v:shape style="width:120pt;height:24.75pt">'
            '<v:imagedata r:id="rId21"/></v:shape></w:pict></w:r>'
        )
    )
    out = tmp_path / name
    doc.save(str(out))
    return out


def test_mixed_text_figure_keeps_paragraph_eid_and_text(tmp_path):
    """The whole point of the pair: a numbered provision that shares its <w:p>
    with an inline formula image keeps its <paragraph> eId and operative text,
    and gains a <figure> — instead of the text vanishing."""
    text = (
        "(a) if the annual pay of the member is less than the prescribed "
        "amount—the number ascertained in accordance with the formula"
    )
    path = _mixed_vml_docx(tmp_path, text)
    meta = _figure_meta("dfrb_act")
    report, xml = _convert_figure_volumes([str(path)], meta, tmp_path)

    root = etree.fromstring(xml.encode())
    ns = {"akn": AKN_NS}
    para_el = root.find(".//akn:paragraph", ns)
    assert para_el is not None, "numbered <paragraph> lost"
    assert para_el.get("eId") == "sec-4A__para-a"
    assert "the number ascertained in accordance with the formula" in "".join(
        para_el.itertext()
    )
    fig = root.find(".//akn:figure", ns)
    assert fig is not None
    assert fig.find("akn:img", ns) is not None
    assert report.figures_found == 1
    # The <figure> nests inside the paragraph it came from, not its predecessor.
    assert para_el.find("akn:figure", ns) is not None


def test_figure_paragraph_text_is_not_dropped_by_builder(meta):
    """Builder-boundary guard: a FIGURE ParsedParagraph that still carries
    non-whitespace text emits that text alongside the <figure>. The reader's
    split normally clears it, so this is the safety net for any other producer."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Rate of pension"),
        ParsedParagraph(
            ElementType.FIGURE,
            text="if the annual pay of the member is less than the prescribed amount",
        ),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    assert xml.find(".//akn:figure", ns) is not None
    assert "if the annual pay of the member" in "".join(xml.itertext())


def test_empty_figure_paragraph_emits_figure_only(meta):
    """Regression: a text-free FIGURE still emits exactly <figure><img/></figure>
    with no stray <p>."""
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Diagrams"),
        ParsedParagraph(ElementType.FIGURE, text="   "),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    fig = xml.find(".//akn:figure", ns)
    assert fig is not None
    assert len(list(fig)) == 1
    assert fig[0].tag == f"{{{AKN_NS}}}img"
    section = xml.find(".//akn:section", ns)
    assert section.findall(".//akn:p", ns) == []


# ---------------------------------------------------------------------------
# Task 13 — XSD family C residue: structural eIds inside quoted content
#
# Re-baselined after Task 6 (`dup-key:eId-act` on part/subsection/section/
# division/chapter/paragraph, filtered from `validate_corpus` against a fresh
# reconvert of the 267 Acts flagged at the 2026-09-07 baseline): Task 6's
# schedule-scoped grouping wrapper and item/quotedStructure nesting resolved
# every cross-item and cross-schedule case, but `_build_quoted_content` itself
# had no collision tracking at all. Two shapes reproduce the residue directly:
#
#   - two structural siblings at the same stack depth inside ONE quoted span
#     (e.g. a duplicated paragraph letter in an inserted list); and
#   - a nested-item quote (`Special ih`/`SubitemHead`) that restarts a whole
#     new numbering sequence more than once inside one span, each restart
#     independently landing on "Part 1" / "Division 1" / etc.
#
# Both mint the identical `@eId` twice today because `_build_quoted_content`
# recomputes each element's eId from (type, num) with no memory of what it
# already emitted.
# ---------------------------------------------------------------------------


def test_quoted_paragraphs_at_same_depth_get_distinct_eids(meta):
    # Two inserted list paragraphs both labelled "(a)" -- a duplicated letter
    # inside one quoted replacement (not two separate items, not two separate
    # quoted spans -- the collision Task 6's item/qstr nesting cannot reach).
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Some Act 2000", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  Section\xa05", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="first thing."),
        ParsedParagraph(ElementType.PARAGRAPH, number="a", text="second thing."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    paras = xml.findall(".//akn:quotedStructure/akn:paragraph", ns)
    assert len(paras) == 2
    assert paras[0].get("eId") != paras[1].get("eId")
    assert [p.find("akn:num", ns).text for p in paras] == ["a", "a"], (
        "the literal letter must still read 'a' in <num> even though the eId is disambiguated"
    )
    # Word-for-word: both paragraphs kept, nothing merged or dropped.
    text = " ".join(" ".join(xml.itertext()).split())
    assert "first thing." in text
    assert "second thing." in text


def test_nested_item_quote_restarting_part_numbering_gets_distinct_eids(meta):
    # One "Insert:" span quotes two whole new nested items (Special ih), each
    # introducing its own "Part 1" -- numbering legitimately restarts per
    # nested item, but nothing inside a quotedStructure tracks that the way
    # the schedule's own grouping wrapper does.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Some Act 2000", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  After Part\xa05", raw_style="ItemHead"),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="5A  Section\xa010", raw_style="Special ih"),
        ParsedParagraph(ElementType.PART, number="1", heading="First new part", raw_style="ActHead 2"),
        ParsedParagraph(ElementType.BODY, text="5B  Section\xa011", raw_style="Special ih"),
        ParsedParagraph(ElementType.PART, number="1", heading="Second new part", raw_style="ActHead 2"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    parts = xml.findall(".//akn:quotedStructure/akn:part", ns)
    assert len(parts) == 2
    assert parts[0].get("eId") != parts[1].get("eId")
    assert [p.find("akn:num", ns).text for p in parts] == ["1", "1"]
    assert [p.find("akn:heading", ns).text for p in parts] == [
        "First new part", "Second new part",
    ]


def test_two_blocklists_in_one_item_across_a_span_get_distinct_eids(meta):
    # Fix-round finding (reviewer): `blocklist_count` resets to 0 at the top
    # of every `_build_quoted_content` call, same as the structural stack
    # does -- two prose chunks in one item that each open a list (split by a
    # quoted span that a margin note then closes) both minted "list-1" under
    # the same `item_eid` prefix. Same collision class as the other two
    # Task 13 tests, in the LIST_ITEM branch instead of the structural one.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.BODY, text="Schedule\xa01—Amendments", raw_style="ActHead 1"),
        ParsedParagraph(ElementType.BODY, text="Some Act 2000", raw_style="ActHead 9"),
        ParsedParagraph(ElementType.BODY, text="1  Section\xa05", raw_style="ItemHead"),
        ParsedParagraph(ElementType.LIST_ITEM, number="1", text="(a) first list item, chunk one."),
        ParsedParagraph(ElementType.BODY, text="Insert:", raw_style="Item"),
        ParsedParagraph(ElementType.BODY, text="something quoted."),
        ParsedParagraph(ElementType.BODY, text="Note: this ends the quoted run", raw_style="note(margin)"),
        ParsedParagraph(ElementType.LIST_ITEM, number="1", text="(a) first list item, chunk two."),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}

    eids = _all_eids(xml)
    assert len(eids) == len(set(eids)), f"eId collision: {sorted(e for e in eids if eids.count(e) > 1)}"

    block_lists = xml.findall(".//akn:hcontainer[@name='item']/akn:blockList", ns)
    assert len(block_lists) == 2
    assert block_lists[0].get("eId") != block_lists[1].get("eId")
    # Word-for-word: both list items kept.
    text = " ".join(" ".join(xml.itertext()).split())
    assert "first list item, chunk one." in text
    assert "first list item, chunk two." in text


def test_quoted_content_disambiguation_does_not_shift_body_eids(meta):
    # Regression gate for the hard constraint: the Task 13 fix lives entirely
    # inside `_build_quoted_content`/`_build_item_body` (schedule-only call
    # graph). A body section elsewhere in the same Act, with no schedule at
    # all, must keep its plain `sec-*`/`subsec-*` eIds -- occurrence-suffixing
    # must never reach a body element.
    paragraphs = [
        ParsedParagraph(ElementType.SECTION, number="1", heading="Short title"),
        ParsedParagraph(ElementType.SECTION, number="5", heading="Interpretation"),
        ParsedParagraph(ElementType.SUBSECTION, number="1", text="In this Act:"),
    ]
    xml, _ = build_xml(meta, paragraphs)
    ns = {"akn": AKN_NS}
    assert xml.find(".//akn:section[@eId='sec-5']", ns) is not None
    assert xml.find(".//akn:subsection[@eId='sec-5__subsec-1']", ns) is not None
